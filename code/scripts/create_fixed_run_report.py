#!/usr/bin/env python3
"""Create a self-contained Word report for a fixed-model threshold rerun."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_table(document: Document, frame: pd.DataFrame) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, frame.columns):
        cell.text = str(value)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = f"{value:.3f}" if isinstance(value, float) else str(value)


def add_bullets(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def metric_view(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame[
        [
            "model", "mae", "rmse", "crps", "outbreak_brier",
            "outbreak_precision", "outbreak_specificity",
            "outbreak_sensitivity", "outbreak_f1",
        ]
    ].copy()
    result.columns = [
        "Model", "MAE", "RMSE", "CRPS", "Brier", "Precision",
        "Specificity", "Sensitivity", "F1",
    ]
    return result


def plot_metrics(summary: pd.DataFrame, output: Path) -> None:
    order = summary.sort_values("mean_metric_rank")["model"].tolist()
    indexed = summary.set_index("model").loc[order]
    colors = [
        "#0072B2" if value == "DeepST-INLA (proposed hybrid)"
        else "#009E73" if value == "DeepST-Alert (proposed deep)"
        else "#A7AFB8"
        for value in order
    ]
    fig, axes = plt.subplots(2, 2, figsize=(14, 10), constrained_layout=True)
    for axis, metric, title in (
        (axes[0, 0], "mae", "Mean H1-H3 MAE (lower is better)"),
        (axes[0, 1], "crps", "Mean H1-H3 CRPS (lower is better)"),
        (axes[1, 0], "outbreak_brier", "Mean H1-H3 Brier score (lower is better)"),
    ):
        axis.barh(order, indexed[metric], color=colors)
        axis.invert_yaxis()
        axis.set_title(title, fontweight="bold")
        axis.grid(axis="x", alpha=0.2)
    axis = axes[1, 1]
    positions = np.arange(len(order))
    width = 0.24
    for offset, metric, label, color in (
        (-width, "outbreak_precision", "Precision", "#E69F00"),
        (0, "outbreak_sensitivity", "Sensitivity", "#009E73"),
        (width, "outbreak_f1", "F1", "#0072B2"),
    ):
        axis.bar(positions + offset, indexed[metric], width, label=label, color=color)
    axis.set_xticks(positions, order, rotation=24, ha="right")
    axis.set_ylim(0, 1)
    axis.set_title("Mean H1-H3 outbreak metrics", fontweight="bold")
    axis.legend(frameon=False)
    axis.grid(axis="y", alpha=0.2)
    fig.savefig(output, dpi=300, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--train-start", required=True)
    parser.add_argument("--train-end", required=True)
    parser.add_argument("--validation-start", required=True)
    parser.add_argument("--validation-end", required=True)
    parser.add_argument("--test-start", required=True)
    parser.add_argument("--test-end", required=True)
    parser.add_argument("--thresholds", type=Path, required=True)
    parser.add_argument("--threshold-column", default="epi_threshold")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    by_horizon = pd.read_csv(args.output / "final_metrics_by_horizon.csv")
    summary = pd.read_csv(args.output / "final_metrics_mean_H1_H3.csv")
    plot_metrics(summary, args.output / "Fig_final_metrics_comparison.png")

    train = f"{args.train_start[:4]}-{args.train_end[:4]}"
    validation = f"{args.validation_start[:4]}-{args.validation_end[:4]}"
    test = f"{args.test_start[:4]}-{args.test_end[:4]}"
    threshold_rows = len(pd.read_csv(args.thresholds, usecols=[args.threshold_column]))
    selection_path = args.run_root / "DEEPST-INLA/selection.json"
    selection = json.loads(selection_path.read_text()) if selection_path.exists() else {}

    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = title.add_run("COMMUNE-LEVEL DENGUE FORECASTING\n")
    heading.bold = True
    heading.font.size = Pt(18)
    subtitle = title.add_run(f"External-threshold rerun; locked test {test}")
    subtitle.italic = True
    subtitle.font.size = Pt(12)

    document.add_heading("1. Evaluation protocol", level=1)
    add_bullets(document, [
        f"Training: {train}; blocked validation and model selection: {validation}; locked test: {test}.",
        "Forecast endpoints are H1, H2, and H3. The quality-controlled panel contains 554 communes.",
        f"Outbreak labels use {args.threshold_column} from {args.thresholds.name} ({threshold_rows:,} commune-month rows). No historical threshold is recalculated by the forecasting pipeline.",
        "The externally supplied table starts in 2015. Earlier targets have no outbreak label and are excluded from the auxiliary alert loss; count-model training is unchanged.",
        "All model architecture, features, optimizer settings, seeds, predictive draws, and validation-only selection rules match the prior final report.",
        "The locked test is not used to choose epochs, ensemble weights, calibration, or alert policy.",
    ])

    document.add_heading("2. Models", level=1)
    add_bullets(document, [
        "INLA-TS: negative-binomial temporal and BYM2 spatial baseline.",
        "TCN-GNN Base: twelve-month causal TCN with four-head graph attention and seven climate variables.",
        "DeepST-Alert: TCN-GNN with a focal-loss auxiliary outbreak head.",
        "TCN-GNN Deep Ensemble: validation-selected stack of TCN-GNN and LSTM-GNN predictive distributions.",
        "Representation-INLA: negative-binomial INLA augmented with residualized deep representation components.",
        "DeepST-INLA: validation-selected multi-objective combination of the deep, alert, and representation-INLA components.",
    ])
    if selection:
        document.add_paragraph(
            f"The recorded selection split is {selection.get('selection_split', validation)}."
        )

    document.add_heading("3. Mean locked-test results, H1-H3", level=1)
    add_table(document, metric_view(summary))
    document.add_picture(
        str(args.output / "Fig_final_metrics_comparison.png"), width=Inches(6.4)
    )

    document.add_heading("4. Horizon-specific results", level=1)
    for horizon in (1, 2, 3):
        document.add_heading(f"H{horizon}", level=2)
        add_table(document, metric_view(by_horizon[by_horizon["horizon"].eq(horizon)]))

    document.add_heading("5. H3 trajectories", level=1)
    document.add_picture(
        str(args.output / "Fig_H3_final_model_comparison.png"), width=Inches(6.4)
    )
    document.add_paragraph(
        "The trajectory panel uses the same four illustrative communes as the prior final report; quantitative conclusions use all test rows."
    )

    document.add_heading("6. Reproducibility note", level=1)
    document.add_paragraph(
        "The report folder contains the complete by-horizon and H1-H3 summary tables. Full predictions, predictive draws, fitted models, training histories, INLA diagnostics, and validation-selection records are retained in the corresponding isolated run folder."
    )
    path = args.output / f"Final_Dengue_Forecasting_Report_thresholds_{train}_{validation}_{test}.docx"
    document.save(path)
    print(path)


if __name__ == "__main__":
    main()
