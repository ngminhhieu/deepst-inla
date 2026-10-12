#!/usr/bin/env python3
"""Build the final English metrics, figures, and Word report."""
from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from multiobjective_report_supplements import (
    add_architecture_and_outbreak_sections,
    add_restricted_test_section,
    add_rolling_refit_section,
    add_table,
    build_architecture_figure,
    build_period_metrics,
    load_rolling_results,
)


ROOT = Path(__file__).resolve().parents[2]
LOSS_COLUMNS = ("mae", "rmse", "crps", "outbreak_brier")
GAIN_COLUMNS = (
    "outbreak_precision", "outbreak_specificity", "outbreak_sensitivity", "outbreak_f1",
)


def sources() -> dict[str, Path]:
    """Locked-test sources; report labels replace internal experiment names."""
    return {
        "INLA-TS": ROOT / "results/eval_2021_2025_v1/INLA-ONLY-test/INLA-TS/metrics.csv",
        "TCN-GNN Base": ROOT / "results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE/metrics.csv",
        "DeepST-Alert (proposed deep)": ROOT / "results/multitask_warning_2021_2025_v1/ALERT-FOCAL-BCE/metrics.csv",
        "TCN-GNN Deep Ensemble (variant)": ROOT / "results/feature_tcn_gnn_2021_2025_v1/ARCH-TOP2-STACK/metrics.csv",
        "Representation-INLA (variant)": ROOT / "results/feature_tcn_gnn_2021_2025_v1/ARCH-WINNER-REPR-INLA/test/combined/metrics.csv",
        "DeepST-INLA (proposed hybrid)": ROOT / "results/final_multiobjective_model_2021_2025/metrics.csv",
    }


def build_metrics(output: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    frames = []
    for label, path in sources().items():
        frame = pd.read_csv(path)
        frame["model"] = label
        frames.append(frame)
    by_horizon = pd.concat(frames, ignore_index=True).sort_values(["model", "horizon"])
    metrics = list(LOSS_COLUMNS + GAIN_COLUMNS)
    summary = by_horizon.groupby("model", as_index=False)[metrics].mean()
    for column in metrics:
        summary[f"rank_{column}"] = summary[column].rank(
            ascending=column in LOSS_COLUMNS, method="min"
        )
    rank_columns = [f"rank_{column}" for column in metrics]
    summary["mean_metric_rank"] = summary[rank_columns].mean(axis=1)
    summary["worst_metric_rank"] = summary[rank_columns].max(axis=1)
    summary = summary.sort_values(["mean_metric_rank", "worst_metric_rank"])
    by_horizon.to_csv(output / "final_metrics_by_horizon.csv", index=False)
    summary.to_csv(output / "final_metrics_mean_H1_H3.csv", index=False)
    return by_horizon, summary


def build_h3_figure(output: Path) -> None:
    figure_base = output / "Fig_H3_final_multiobjective"
    command = [
        sys.executable,
        str(ROOT / "code/scripts/make_comparison_h3_outbreak_figure.py"),
        "--results-root", str(ROOT),
        "--output", str(figure_base),
        "--communes", "28159,27976,31147,31168",
        "--models",
        f"inla={ROOT / 'results/eval_2021_2025_v1/INLA-ONLY-test/INLA-TS/predictions.csv'}",
        f"base={ROOT / 'results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE/predictions.csv'}",
        f"representation={ROOT / 'results/feature_tcn_gnn_2021_2025_v1/ARCH-WINNER-REPR-INLA/test/combined/predictions.csv'}",
        f"hybrid={ROOT / 'results/final_multiobjective_model_2021_2025/predictions.csv'}",
        f"deep={ROOT / 'results/multitask_warning_2021_2025_v1/ALERT-FOCAL-BCE/predictions.csv'}",
        "--labels",
        "INLA-TS",
        "TCN-GNN Base",
        "Representation-INLA",
        "DeepST-INLA (proposed hybrid)",
        "DeepST-Alert (proposed deep)",
        "--title",
        "Locked 2021-2025 H3 comparison on the original illustrative communes",
    ]
    subprocess.run(command, check=True)
    figure_base.with_name(figure_base.name + "_selection.csv").unlink(missing_ok=True)


def build_restricted_h3_figure(output: Path) -> None:
    figure_base = output / "Fig_H3_final_multiobjective_2023_2025"
    command = [
        sys.executable,
        str(ROOT / "code/scripts/make_comparison_h3_outbreak_figure.py"),
        "--results-root", str(ROOT),
        "--output", str(figure_base),
        "--communes", "28159,27976,31147,31168",
        "--start-date", "2023-01-01",
        "--end-date", "2025-12-01",
        "--models",
        f"inla={ROOT / 'results/eval_2021_2025_v1/INLA-ONLY-test/INLA-TS/predictions.csv'}",
        f"base={ROOT / 'results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE/predictions.csv'}",
        f"representation={ROOT / 'results/feature_tcn_gnn_2021_2025_v1/ARCH-WINNER-REPR-INLA/test/combined/predictions.csv'}",
        f"hybrid={ROOT / 'results/final_multiobjective_model_2021_2025/predictions.csv'}",
        f"deep={ROOT / 'results/multitask_warning_2021_2025_v1/ALERT-FOCAL-BCE/predictions.csv'}",
        "--labels",
        "INLA-TS",
        "TCN-GNN Base",
        "Representation-INLA",
        "DeepST-INLA (proposed hybrid)",
        "DeepST-Alert (proposed deep)",
        "--title",
        "Restricted 2023-2025 H3 comparison on the original illustrative communes",
    ]
    subprocess.run(command, check=True)
    figure_base.with_name(figure_base.name + "_selection.csv").unlink(missing_ok=True)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_pseudocode(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="No Spacing")
    for index, line in enumerate(lines):
        run = paragraph.add_run(("\n" if index else "") + line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)


def add_labelled_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(label + ". ")
    run.bold = True
    paragraph.add_run(text)


def display_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame[[
        "model", "mae", "rmse", "crps", "outbreak_brier", "outbreak_precision",
        "outbreak_specificity", "outbreak_sensitivity", "outbreak_f1",
    ]].copy()
    output.columns = [
        "Model", "MAE", "RMSE", "CRPS", "Brier", "Precision", "Specificity",
        "Sensitivity", "F1",
    ]
    return output


def write_report(
    output: Path,
    by_horizon: pd.DataFrame,
    summary: pd.DataFrame,
    rolling_results: tuple[pd.DataFrame, pd.DataFrame, dict[str, object]],
    restricted_results: tuple[pd.DataFrame, pd.DataFrame],
) -> Path:
    """Write the complete publication-facing English report."""
    feature = pd.read_csv(
        ROOT / "results/feature_ablation_selected_base_2021_2025_v1/feature_ablation_validation_summary.csv"
    )
    architecture = pd.read_csv(
        ROOT / "results/feature_tcn_gnn_2021_2025_v1/architecture_validation_summary.csv"
    )

    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(24)
    document.styles["Heading 2"].font.size = Pt(12.5)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("COMMUNE-LEVEL DENGUE FORECASTING")
    run.bold = True
    run.font.size = Pt(19)
    line_break = title.add_run("\n")
    line_break.bold = True
    line_break.font.size = Pt(19)
    subtitle = title.add_run("Model development and test evaluation, 2021–2025")
    subtitle.italic = True
    subtitle.font.size = Pt(12)

    document.add_heading("1. Evaluation protocol", level=1)
    add_bullets(document, [
        "Data split: training 2010–2018; blocked validation 2019–2020; test 2021–2025.",
        "Forecast endpoints: H1, H2, and H3. Each horizon contains 60 months × 554 communes = 33,240 test predictions.",
        "Quality control: the source contained 555 communes. Fcode 31495 (Xa Phien) was excluded because the recorded population (47–50), density (0.50 persons/km²), and maximum incidence (36,000 per 100,000) were implausible. The final panel contains 554 communes.",
        "Outbreak definition: observed cases exceed the commune- and target-month-specific mean + 2 sample standard deviations from the same calendar month in the previous five years.",
    ])

    document.add_heading("1.1 Evaluation metrics", level=2)
    document.add_paragraph()
    metric_table = pd.DataFrame([
        ("MAE", "Lower", "Mean absolute incidence error"),
        ("RMSE", "Lower", "Incidence error with stronger penalty for large misses"),
        ("CRPS", "Lower", "Accuracy and calibration of the predictive distribution"),
        ("Brier", "Lower", "Error of predicted outbreak probabilities"),
        ("Precision", "Higher", "Fraction of alerts that are correct"),
        ("Specificity", "Higher", "Fraction of non-outbreaks without a false alert"),
        ("Sensitivity/Recall", "Higher", "Fraction of observed outbreaks detected"),
        ("F1", "Higher", "Balance between precision and sensitivity"),
    ], columns=["Metric", "Better", "Meaning"])
    add_table(document, metric_table, widths=(816_610, 730_250, 5_105_400))

    document.add_heading("1.2 Output required from another model", level=2)
    document.add_paragraph(
        "A model trained elsewhere can be evaluated by returning predictions for every test key. "
        "Ground truth, population, and outbreak thresholds are added from the common evaluation template."
    )
    document.add_paragraph()
    schema = pd.DataFrame([
        ("28159", "2021-09-01", "2021-10-01", "1", "35.2", "17.4", "61.8", "0.31", "0.33", "draws_000001"),
        ("28159", "2021-09-01", "2021-11-01", "2", "39.7", "18.9", "70.5", "0.42", "0.35", "draws_000002"),
        ("28159", "2021-09-01", "2021-12-01", "3", "42.6", "18.1", "77.4", "0.63", "0.37", "draws_000003"),
        ("27976", "2021-09-01", "2021-10-01", "1", "51.3", "24.8", "89.6", "0.48", "0.33", "draws_000004"),
        ("27976", "2021-09-01", "2021-11-01", "2", "58.9", "27.2", "103.1", "0.57", "0.35", "draws_000005"),
        ("27976", "2021-09-01", "2021-12-01", "3", "64.5", "29.7", "116.8", "0.71", "0.37", "draws_000006"),
    ], columns=[
        "fcode", "issue_date", "target_date", "horizon", "predicted_mean",
        "lower_80", "upper_80", "outbreak_probability",
        "outbreak_decision_threshold", "draws_row_id",
    ])
    add_table(
        document,
        schema,
        widths=(
            441_325,
            585_470,
            602_615,
            534_035,
            701_040,
            563_245,
            566_420,
            867_410,
            1_144_905,
            645_795,
        ),
    )
    document.add_paragraph(
        "Each draws_row_id points to 500 predictive draws in a separate array. The values above are "
        "illustrative; submitted files must cover every test row."
    )

    document.add_heading("2. Base architecture ablation", level=1)
    document.add_paragraph(
        "All candidates used the full feature set. TCN-GNN achieved the best validation CRPS and F1 "
        "and was selected as the base architecture."
    )
    add_labelled_paragraph(
        document,
        "Last-state Graph Neural Network (GNN)",
        "Uses only the most recent monthly state as input. The GNN learns spatial dependencies between "
        "communes without explicitly modelling temporal patterns across previous months.",
    )
    add_labelled_paragraph(
        document,
        "TCN-GNN",
        "Uses a Temporal Convolutional Network (TCN) to learn temporal patterns across historical "
        "observations, followed by the GNN to model spatial dependencies between communes. This is "
        "selected as the base model.",
    )
    add_labelled_paragraph(
        document,
        "LSTM-GNN",
        "Uses Long Short-Term Memory (LSTM) to learn temporal dependencies across previous months, "
        "followed by the GNN to capture spatial relationships between communes.",
    )
    add_labelled_paragraph(
        document,
        "CNN-LSTM-GNN",
        "Uses a Convolutional Neural Network (CNN) to extract local temporal patterns, followed by LSTM "
        "to capture longer-term temporal dependencies and the GNN to model spatial relationships between "
        "communes.",
    )
    add_labelled_paragraph(
        document,
        "TCN-GNN with dynamic spatial gate",
        "Extends TCN-GNN with a learnable sigmoid gate that controls how much neighbouring information "
        "is incorporated into each commune’s local representation.",
    )
    architecture_table = architecture[[
        "model", "mean_crps", "mean_mae", "mean_brier", "mean_sensitivity", "mean_f1"
    ]].copy()
    architecture_table["model"] = architecture_table["model"].replace({
        "ARCH-LAST-GNN": "Last-state GNN",
        "ARCH-TCN-GNN-NOGATE": "TCN-GNN",
        "ARCH-TCN-GNN-GATED": "TCN-GNN + dynamic spatial gate",
        "ARCH-LSTM-GNN": "LSTM-GNN",
        "ARCH-CNN-LSTM-GNN": "CNN-LSTM-GNN",
    })
    architecture_table.columns = ["Architecture", "CRPS", "MAE", "Brier", "Sensitivity", "F1"]
    add_table(
        document,
        architecture_table,
        widths=(1_028_700, 1_097_280, 1_097_280, 1_097_280, 1_097_280, 1_097_280),
    )
    document.add_paragraph()

    document.add_heading("3. Feature ablation on the selected base", level=1)
    document.add_paragraph(
        "Feature ablation was rerun using the selected TCN-GNN. The Basic climate set contains "
        "temperature, relative humidity, and rainfall. Each “Basic +” row adds one additional variable "
        "to this same Basic climate set; the additions are not cumulative. The Full set combines Basic "
        "climate with all four additional variables. The Full set achieved the best CRPS, Brier score, "
        "sensitivity, and F1 and was therefore selected for the final models."
    )
    add_labelled_paragraph(
        document, "Basic climate",
        "Temperature, relative humidity, and rainfall describe heat, atmospheric moisture, and water input. "
        "Source: the eDengue commune climate collections.",
    )
    add_labelled_paragraph(
        document, "Soil moisture",
        "Represents persistence of wet breeding habitat after rainfall. Source: NASA POWER GWETTOP.",
    )
    add_labelled_paragraph(
        document, "Solar radiation",
        "Represents surface energy affecting temperature and mosquito development. Source: NASA POWER "
        "ALLSKY_SFC_SW_DWN.",
    )
    add_labelled_paragraph(
        document, "Dew point",
        "Represents absolute atmospheric moisture more directly than relative humidity alone. Source: "
        "NASA POWER T2MDEW.",
    )
    document.add_paragraph(
        "ET0 (reference evapotranspiration). Represents evaporative demand and drying conditions. "
        "Source: Open-Meteo Historical Weather API / ERA5-Land FAO reference evapotranspiration."
    )
    document.add_paragraph(
        "Full set. Combines the Basic climate variables with soil moisture, solar radiation, dew point, "
        "and ET0. This is the selected feature set."
    )
    feature_table = feature[[
        "model", "mean_crps", "mean_mae", "mean_brier", "mean_sensitivity", "mean_f1"
    ]].copy()
    feature_table["model"] = feature_table["model"].replace({
        "TCN-GNN-BASE": "Basic climate\n(temperature,\nhumidity, rainfall)",
        "TCN-GNN-SOIL": "Basic + soil\nmoisture",
        "TCN-GNN-SOLAR": "Basic + solar\nradiation",
        "TCN-GNN-DEWPOINT": "Basic + dew point",
        "TCN-GNN-ET0": "Basic + ET0",
        "TCN-GNN-ALL": "Full set (Basic + all\nfour additions)",
    })
    feature_order = [
        "Basic climate\n(temperature,\nhumidity, rainfall)",
        "Basic + ET0",
        "Basic + soil\nmoisture",
        "Basic + solar\nradiation",
        "Basic + dew point",
        "Full set (Basic + all\nfour additions)",
    ]
    feature_table["model"] = pd.Categorical(
        feature_table["model"], categories=feature_order, ordered=True
    )
    feature_table = feature_table.sort_values("model")
    feature_table["model"] = feature_table["model"].astype(str)
    feature_table.columns = ["Feature set", "CRPS", "MAE", "Brier", "Sensitivity", "F1"]
    add_table(document, feature_table)
    document.add_paragraph()

    document.add_heading("4. Final models", level=1)
    document.add_heading("4.1 DeepST-INLA: selected final model", level=2)
    document.add_paragraph(
        "DeepST-INLA uses the selected TCN-GNN to learn nonlinear temporal patterns from historical dengue "
        "observations and spatial relationships between communes. It produces both a predictive case-count "
        "distribution and a direct outbreak probability. The outbreak objective uses focal binary "
        "cross-entropy (BCE), which places more emphasis on difficult and relatively rare outbreak cases."
    )
    document.add_paragraph(
        "The learned deep representation is then passed to Integrated Nested Laplace Approximation (INLA). "
        "INLA provides a structured probabilistic count model with explicit temporal and spatial effects "
        "and posterior uncertainty. Validation is used to determine how the deep forecast should support "
        "the INLA distribution in elevated-risk forecasts. The final outbreak probability is then "
        "calculated from the count distribution and the direct outbreak signal using the common procedure "
        "in Section 9."
    )
    document.add_paragraph("Pseudocode:")
    document.add_paragraph(
        "The pseudocode emphasises the purpose of each step rather than implementation details. Parameters "
        "controlling the INLA/deep combination and the final alert calculation are selected on validation data."
    )
    add_pseudocode(document, [
        "1  h, D_deep, p_aux = TCN_GNN(x)",
        "2  loss = CountNLL(D_deep,y) + 0.25*FocalBCE(p_aux,label)",
        "3  z = PrepareForINLA(h)",
        "4  D_inla = INLA(z,covariates,BYM2,500)",
        "5  D_final = INLASupport(D_inla,D_deep,p_aux,h)   # validation-selected",
        "6  p_count = Mean(D_final > outbreak_threshold)",
        "7  p_raw = w_count[h]*p_count + w_aux[h]*p_aux",
        "8  p_final = Platt_h(p_raw)",
        "9  alert = (p_final >= alert_threshold[h])",
    ])
    document.add_paragraph(
        "The key role of INLA is therefore to provide the structured probabilistic count forecast. The "
        "deep model supplies nonlinear temporal-spatial information and an outbreak-focused signal, while "
        "validation determines how these components are combined without using the test period."
    )

    document.add_heading("4.2 DeepST-Alert: deep-only comparator", level=2)
    document.add_paragraph(
        "DeepST-Alert is the deep-learning-only comparator. It uses the same TCN-GNN inputs, the same "
        "count-forecasting objective, the same focal-BCE outbreak objective, and the same "
        "outbreak-probability calculation as DeepST-INLA. The intended difference is that DeepST-Alert "
        "does not use INLA: its predictive case-count distribution comes directly from the deep model."
    )
    document.add_paragraph(
        "This comparison isolates the value of adding the structured INLA component. Both models still "
        "estimate outbreak risk from two sources—the predicted count distribution and the direct outbreak "
        "probability—and both combine and calibrate those probabilities using parameters selected on "
        "validation data."
    )
    document.add_paragraph("Pseudocode:")
    add_pseudocode(document, [
        "1  h, D_deep, p_aux = TCN_GNN(x)",
        "2  loss = CountNLL(D_deep,y) + 0.25*FocalBCE(p_aux,label)",
        "3  D_final = D_deep                              # no INLA",
        "4  p_count = Mean(D_final > outbreak_threshold)",
        "5  p_raw = w_count[h]*p_count + w_aux[h]*p_aux",
        "6  p_final = Platt_h(p_raw)",
        "7  alert = (p_final >= alert_threshold[h])",
    ])
    document.add_paragraph(
        "The count distribution is still required because it provides the case forecast, uncertainty "
        "intervals, and the count-based outbreak probability p_count. The direct probability p_aux "
        "provides a complementary signal trained specifically to recognise outbreak cases."
    )

    indexed = summary.set_index("model")
    deep = indexed.loc["DeepST-Alert (proposed deep)"]
    hybrid = indexed.loc["DeepST-INLA (proposed hybrid)"]
    document.add_paragraph(
        f"DeepST-INLA was selected because, compared with DeepST-Alert, it reduced mean Brier score from "
        f"{deep.outbreak_brier:.3f} to {hybrid.outbreak_brier:.3f} and increased F1 from "
        f"{deep.outbreak_f1:.3f} to {hybrid.outbreak_f1:.3f}. Sensitivity decreased from "
        f"{deep.outbreak_sensitivity:.3f} to {hybrid.outbreak_sensitivity:.3f} because the hybrid removed "
        "many false alerts; the selection therefore reflects overall balance rather than maximising "
        "sensitivity alone."
    )

    document.add_heading("5. Test results", level=1)
    document.add_paragraph()
    add_table(document, display_metrics(summary))
    document.add_paragraph(
        "DeepST-INLA provides the best overall balance. Representation-INLA and the deep ensemble are "
        "intermediate variants of the proposed hybrid and proposed deep model, respectively."
    )

    document.add_heading("5.1 H3 visualization", level=2)
    document.add_paragraph()
    h3_order = [
        "INLA-TS", "TCN-GNN Base", "Representation-INLA (variant)",
        "DeepST-INLA (proposed hybrid)", "DeepST-Alert (proposed deep)",
    ]
    h3 = by_horizon[by_horizon["horizon"].eq(3)].set_index("model").loc[h3_order].reset_index()
    add_table(document, display_metrics(h3))
    document.add_picture(str(output / "Fig_H3_final_multiobjective.png"), width=Inches(7.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "The H3 figure retains the four illustrative communes 28159, 27976, 31147, and 31168."
    )

    document.add_heading("6. Conclusion", level=1)
    conclusion = document.add_paragraph()
    selected = conclusion.add_run("Selected final model: DeepST-INLA. ")
    selected.bold = True
    conclusion.add_run(
        "It is retained because it gives the strongest overall trade-off among count accuracy, "
        "predictive-distribution quality, probability calibration, false-alert control, and outbreak "
        "detection. The detailed architecture and probability calculation are documented in Sections 8 "
        "and 9 so that the final implementation can be reproduced without relying on shorthand model names."
    )

    rolling_metrics, rolling_comparison, rolling_manifest = rolling_results
    add_rolling_refit_section(
        document,
        rolling_metrics,
        rolling_comparison,
        rolling_manifest,
    )
    add_architecture_and_outbreak_sections(
        document,
        output / "Fig_DeepST_INLA_architecture.png",
        ROOT / "results/final_multiobjective_model_2021_2025/selection.json",
    )
    restricted_by_horizon, restricted_summary = restricted_results
    add_restricted_test_section(
        document,
        restricted_by_horizon,
        restricted_summary,
        output / "Fig_H3_final_multiobjective_2023_2025.png",
    )

    path = output / "Final_Dengue_Forecasting_Report_2021_2025.docx"
    document.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output", type=Path,
        default=ROOT / "results/current",
    )
    parser.add_argument(
        "--rolling-results",
        type=Path,
        default=ROOT
        / "results/final_multiobjective_rolling_refit_2021_2025"
        / "aggregated_test",
    )
    args = parser.parse_args()
    rolling_results = load_rolling_results(args.rolling_results)
    args.output.mkdir(parents=True, exist_ok=True)
    by_horizon, summary = build_metrics(args.output)
    restricted_results = build_period_metrics(
        {model: metrics_path.parent for model, metrics_path in sources().items()},
        args.output,
    )
    build_h3_figure(args.output)
    build_restricted_h3_figure(args.output)
    build_architecture_figure(args.output / "Fig_DeepST_INLA_architecture.png")
    write_report(
        args.output,
        by_horizon,
        summary,
        rolling_results,
        restricted_results,
    )


if __name__ == "__main__":
    main()
