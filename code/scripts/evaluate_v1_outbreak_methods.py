#!/usr/bin/env python3
"""Evaluate the three legacy v1 outbreak rules on saved model draws.

The incidence cutoff for the two consensus rules is selected on validation
and then applied unchanged to the locked test. No model is fitted here.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


INCIDENCE_CANDIDATES = (0, 10, 15, 20, 25, 30, 35, 40, 45, 50)
METHOD_LABELS = {
    "current_probability": "Current calibrated probability",
    "risk_abs": "v1 absolute-risk threshold",
    "risk_z": "v1 z-score threshold",
    "pred_mean": "Forecast mean > outbreak threshold",
    "two_of_three": "v1 consensus: at least 2 of 3",
    "three_of_three": "v1 consensus: all 3",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--validation-predictions", type=Path, required=True)
    parser.add_argument("--validation-samples", type=Path, required=True)
    parser.add_argument("--test-predictions", type=Path, required=True)
    parser.add_argument("--test-samples", type=Path, required=True)
    parser.add_argument("--monthly-data", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--title", default="Three v1 outbreak methods")
    parser.add_argument(
        "--run-mode", choices=("no_retrain", "retrained"), default="no_retrain"
    )
    return parser.parse_args()


def load_samples(path: Path) -> np.ndarray:
    archive = np.load(path)
    key = "samples" if "samples" in archive.files else archive.files[0]
    values = archive[key]
    if values.ndim != 4:
        raise ValueError(f"Expected a four-dimensional sample cube, got {values.shape}")
    return values


def load_predictions(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype={"fcode": str})
    frame["target_date"] = pd.to_datetime(frame["target_date"])
    frame["issue_date"] = pd.to_datetime(frame["issue_date"])
    required = {
        "fcode", "target_date", "issue_date", "horizon", "target_cases",
        "population", "predicted_mean", "outbreak_threshold",
        "outbreak_probability", "outbreak_decision_threshold",
    }
    missing = sorted(required.difference(frame.columns))
    if missing:
        raise ValueError(f"Prediction file is missing columns: {missing}")
    return frame


def historical_moments(frame: pd.DataFrame, monthly_path: Path) -> tuple[np.ndarray, np.ndarray]:
    monthly = pd.read_csv(monthly_path, dtype={"fcode": str})
    monthly["date"] = pd.to_datetime(
        {"year": monthly["year"], "month": monthly["month"], "day": 1}
    )
    lookup = monthly.set_index(["fcode", "date"])["obs_dengue_cases"]
    means = np.empty(len(frame), dtype=float)
    sds = np.empty(len(frame), dtype=float)
    for index, row in enumerate(frame[["fcode", "target_date"]].itertuples(index=False)):
        values = []
        for lag in range(1, 6):
            key = (row.fcode, row.target_date - pd.DateOffset(years=lag))
            value = lookup.get(key, np.nan)
            if pd.notna(value):
                values.append(float(value))
        if len(values) < 2:
            means[index], sds[index] = np.nan, np.nan
        else:
            means[index] = float(np.mean(values))
            sds[index] = float(np.std(values, ddof=1))
    return means, sds


def selected_risk_thresholds(samples: np.ndarray, mean: np.ndarray, sd: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Reproduce the two v1 argmax risk thresholds for each forecast row."""
    integer_samples = np.asarray(samples, dtype=np.int64)
    absolute = np.full(len(integer_samples), np.nan, dtype=float)
    zscore = np.full(len(integer_samples), np.nan, dtype=float)
    for index, forecast in enumerate(integer_samples):
        maximum = int(forecast.max(initial=0))
        if maximum <= 0:
            continue
        thresholds = np.arange(1, maximum + 1, dtype=np.int64)
        ordered = np.sort(forecast)
        exceedance = (
            len(forecast) - np.searchsorted(ordered, thresholds, side="right")
        ) / len(forecast)
        absolute[index] = float(thresholds[int(np.argmax(thresholds * exceedance))])
        if np.isfinite(sd[index]) and sd[index] > 0:
            score = exceedance * (thresholds - mean[index]) / sd[index]
            zscore[index] = float(thresholds[int(np.argmax(score))])
    return absolute, zscore


def add_flags(frame: pd.DataFrame, cube: np.ndarray, monthly_path: Path) -> pd.DataFrame:
    output = frame.copy()
    output["observed_outbreak"] = output["target_cases"] > output["outbreak_threshold"]
    output["pred_mean_incidence"] = output["predicted_mean"] / output["population"] * 100_000
    output["historical_mean"], output["historical_sd"] = historical_moments(output, monthly_path)
    reconstructed = output["historical_mean"] + 2 * output["historical_sd"]
    difference = np.abs(reconstructed - output["outbreak_threshold"])
    if float(np.nanmax(difference)) > 2e-3:
        raise ValueError(
            "Saved outbreak thresholds do not match the five-year mean + 2 SD definition; "
            f"maximum difference={float(np.nanmax(difference)):.6g}"
        )

    output["flag_current_probability"] = (
        output["outbreak_probability"] >= output["outbreak_decision_threshold"]
    )
    output["risk_abs_threshold"] = np.nan
    output["risk_z_threshold"] = np.nan
    output["flag_risk_abs"] = False
    output["flag_risk_z"] = False
    output["flag_pred_mean"] = output["predicted_mean"] > output["outbreak_threshold"]
    horizons = sorted(int(value) for value in output["horizon"].unique())
    if cube.shape[2] != len(horizons):
        raise ValueError(f"Sample horizons {cube.shape} do not match {horizons}")
    for horizon_index, horizon in enumerate(horizons):
        mask = output["horizon"].eq(horizon)
        count = int(mask.sum())
        values = cube[:, :, horizon_index, :].reshape(-1, cube.shape[-1])
        if len(values) != count:
            raise ValueError(f"H{horizon}: {len(values)} sample rows != {count} predictions")
        part = output.loc[mask]
        absolute, zscore = selected_risk_thresholds(
            values,
            part["historical_mean"].to_numpy(float),
            part["historical_sd"].to_numpy(float),
        )
        output.loc[mask, "risk_abs_threshold"] = absolute
        output.loc[mask, "risk_z_threshold"] = zscore
        output.loc[mask, "flag_risk_abs"] = absolute > part["outbreak_threshold"].to_numpy(float)
        output.loc[mask, "flag_risk_z"] = (
            np.isfinite(zscore) & (zscore > part["outbreak_threshold"].to_numpy(float))
        )
    output["methods_count"] = output[["flag_risk_abs", "flag_risk_z", "flag_pred_mean"]].sum(axis=1)
    return output


def binary_metrics(observed: np.ndarray, predicted: np.ndarray) -> dict[str, float | int]:
    observed = np.asarray(observed, dtype=bool)
    predicted = np.asarray(predicted, dtype=bool)
    tp = int(np.sum(observed & predicted)); fp = int(np.sum(~observed & predicted))
    fn = int(np.sum(observed & ~predicted)); tn = int(np.sum(~observed & ~predicted))
    ratio = lambda a, b: float(a / b) if b else 0.0
    precision = ratio(tp, tp + fp); sensitivity = ratio(tp, tp + fn)
    specificity = ratio(tn, tn + fp)
    return {
        "n": len(observed), "observed_positive": int(observed.sum()),
        "predicted_positive": int(predicted.sum()), "tp": tp, "fp": fp,
        "fn": fn, "tn": tn, "accuracy": ratio(tp + tn, len(observed)),
        "precision": precision, "specificity": specificity,
        "sensitivity": sensitivity,
        "balanced_accuracy": (sensitivity + specificity) / 2,
        "f1": ratio(2 * precision * sensitivity, precision + sensitivity),
    }


def evaluate_grid(frame: pd.DataFrame, split: str) -> pd.DataFrame:
    rows = []
    groups = [(h, frame[frame["horizon"].eq(h)]) for h in sorted(frame["horizon"].unique())]
    groups.append(("all", frame))
    for horizon, part in groups:
        observed = part["observed_outbreak"].to_numpy(bool)
        for method, column in (
            ("current_probability", "flag_current_probability"),
            ("risk_abs", "flag_risk_abs"), ("risk_z", "flag_risk_z"),
            ("pred_mean", "flag_pred_mean"),
        ):
            rows.append({"split": split, "horizon": horizon, "method": method,
                         "incidence_cutoff": np.nan,
                         **binary_metrics(observed, part[column].to_numpy(bool))})
        for cutoff in INCIDENCE_CANDIDATES:
            incidence_ok = part["pred_mean_incidence"].to_numpy(float) > cutoff
            for method, count in (("two_of_three", 2), ("three_of_three", 3)):
                predicted = (part["methods_count"].to_numpy(int) >= count) & incidence_ok
                rows.append({"split": split, "horizon": horizon, "method": method,
                             "incidence_cutoff": cutoff,
                             **binary_metrics(observed, predicted)})
    return pd.DataFrame(rows)


def select_policies(validation_metrics: pd.DataFrame) -> pd.DataFrame:
    selected = []
    numeric = validation_metrics[validation_metrics["horizon"].ne("all")].copy()
    numeric["horizon"] = numeric["horizon"].astype(int)
    for horizon in sorted(numeric["horizon"].unique()):
        for method in ("two_of_three", "three_of_three"):
            candidates = numeric[(numeric["horizon"].eq(horizon)) & (numeric["method"].eq(method))]
            winner = candidates.sort_values(
                ["f1", "balanced_accuracy", "specificity", "incidence_cutoff"],
                ascending=[False, False, False, True],
            ).iloc[0]
            selected.append({"horizon": horizon, "method": method,
                             "incidence_cutoff": int(winner["incidence_cutoff"]),
                             "validation_f1": winner["f1"],
                             "validation_balanced_accuracy": winner["balanced_accuracy"]})
    return pd.DataFrame(selected)


def selected_test_metrics(test: pd.DataFrame, policies: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for horizon in sorted(test["horizon"].unique()):
        part = test[test["horizon"].eq(horizon)]
        observed = part["observed_outbreak"].to_numpy(bool)
        for method, column in (
            ("current_probability", "flag_current_probability"),
            ("risk_abs", "flag_risk_abs"), ("risk_z", "flag_risk_z"),
            ("pred_mean", "flag_pred_mean"),
        ):
            rows.append({"horizon": int(horizon), "method": method,
                         "incidence_cutoff": np.nan,
                         **binary_metrics(observed, part[column].to_numpy(bool))})
        for method, count in (("two_of_three", 2), ("three_of_three", 3)):
            cutoff = int(policies.loc[
                policies["horizon"].eq(int(horizon)) & policies["method"].eq(method),
                "incidence_cutoff",
            ].iloc[0])
            predicted = ((part["methods_count"].to_numpy(int) >= count)
                         & (part["pred_mean_incidence"].to_numpy(float) > cutoff))
            rows.append({"horizon": int(horizon), "method": method,
                         "incidence_cutoff": cutoff, **binary_metrics(observed, predicted)})
    return pd.DataFrame(rows)


def plot_metrics(metrics: pd.DataFrame, path: Path) -> None:
    summary = metrics.groupby("method", as_index=False)[["precision", "specificity", "sensitivity", "f1"]].mean()
    order = summary.sort_values("f1", ascending=False)["method"].tolist()
    values = summary.set_index("method").loc[order]
    x = np.arange(len(order)); width = 0.19
    fig, axis = plt.subplots(figsize=(12, 6), constrained_layout=True)
    for offset, column, label in ((-1.5, "precision", "Precision"), (-.5, "specificity", "Specificity"),
                                  (.5, "sensitivity", "Sensitivity"), (1.5, "f1", "F1")):
        axis.bar(x + offset * width, values[column], width, label=label)
    axis.set_xticks(x, [METHOD_LABELS[m] for m in order], rotation=18, ha="right")
    axis.set_ylim(0, 1); axis.grid(axis="y", alpha=.25); axis.legend(frameon=False)
    axis.set_title("Locked test 2021–2025: outbreak method comparison")
    fig.savefig(path, dpi=220); plt.close(fig)


def add_table(document: Document, frame: pd.DataFrame) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns)); table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, frame.columns): cell.text = str(column)
    for values in frame.itertuples(index=False, name=None):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = f"{value:.3f}" if isinstance(value, (float, np.floating)) else str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs: run.font.size = Pt(8)


def write_report(
    title: str, metrics: pd.DataFrame, policies: pd.DataFrame, output: Path, run_mode: str
) -> None:
    mean = metrics.groupby("method", as_index=False)[["precision", "specificity", "sensitivity", "balanced_accuracy", "f1"]].mean()
    mean["method"] = mean["method"].map(METHOD_LABELS)
    mean = mean.sort_values("f1", ascending=False)
    best = mean.iloc[0]
    current = mean[mean["method"].eq(METHOD_LABELS["current_probability"])].iloc[0]
    document = Document(); document.styles["Normal"].font.name = "Calibri"
    title_p = document.add_paragraph(); title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title); run.bold = True; run.font.size = Pt(16)
    document.add_heading("Protocol", level=1)
    lead = (
        "No model was retrained; existing predictive draws were post-processed."
        if run_mode == "no_retrain"
        else "The existing DeepST-INLA pipeline was freshly retrained without architecture search."
    )
    document.add_paragraph(
        lead + " The fixed protocol uses 2010–2018 training, "
        "2019–2020 rolling issue-date validation, refitting on 2010–2020, and locked 2021–2025 test. The outbreak "
        "limit is the existing commune/month five-year mean + 2 SD threshold. The three legacy v1 "
        "signals are absolute-risk, z-score-risk, and forecast-mean exceedance. Incidence cutoffs for "
        "the 2-of-3 and 3-of-3 rules were selected only on validation and then locked for test."
    )
    document.add_heading("Validation-selected incidence cutoffs", level=1)
    add_table(document, policies)
    document.add_heading("Locked-test results", level=1)
    add_table(document, mean)
    document.add_paragraph(
        f"The highest mean H1–H3 F1 was {best['method']} ({best['f1']:.3f}). "
        f"The current calibrated-probability rule achieved {current['f1']:.3f}. "
        "This comparison changes only the alert decision layer; count forecasts and predictive draws are identical."
    )
    document.add_picture(str(output / "outbreak_method_comparison.png"), width=Inches(6.4))
    suffix = "No_Retrain" if run_mode == "no_retrain" else "After_Retraining"
    document.save(output / f"Outbreak_Method_Comparison_{suffix}.docx")

    mean_doc = Document()
    mean_doc.add_heading(f"{title}: forecast-mean rule", level=0)
    mean_doc.add_paragraph(
        "This report isolates the deterministic rule predicted_mean > commune/month outbreak threshold. "
        "It does not use voting or an incidence cutoff."
    )
    add_table(mean_doc, metrics[metrics["method"].isin(["pred_mean", "current_probability"])][
        ["horizon", "method", "precision", "specificity", "sensitivity", "balanced_accuracy", "f1"]
    ])
    mean_doc.save(output / f"Mean_Threshold_Report_{suffix}.docx")

    three_doc = Document()
    three_doc.add_heading(f"{title}: three v1 methods and consensus", level=0)
    three_doc.add_paragraph(
        "This report contains the absolute-risk, z-score-risk, and forecast-mean signals, plus validation-selected "
        "2-of-3 and 3-of-3 consensus rules. Incidence cutoffs are selected only on validation."
    )
    add_table(three_doc, policies)
    add_table(three_doc, metrics[metrics["method"].ne("current_probability")][
        ["horizon", "method", "incidence_cutoff", "precision", "specificity", "sensitivity", "balanced_accuracy", "f1"]
    ])
    three_doc.save(output / f"Three_Methods_Report_{suffix}.docx")


def main() -> None:
    args = parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    validation = add_flags(load_predictions(args.validation_predictions), load_samples(args.validation_samples), args.monthly_data)
    test = add_flags(load_predictions(args.test_predictions), load_samples(args.test_samples), args.monthly_data)
    validation_metrics = evaluate_grid(validation, "validation")
    test_grid = evaluate_grid(test, "test")
    policies = select_policies(validation_metrics)
    test_selected = selected_test_metrics(test, policies)
    validation_metrics.to_csv(args.output / "validation_method_grid.csv", index=False)
    test_grid.to_csv(args.output / "test_method_grid.csv", index=False)
    policies.to_csv(args.output / "selected_incidence_cutoffs.csv", index=False)
    test_selected.to_csv(args.output / "test_selected_method_metrics.csv", index=False)
    keep = ["fcode", "issue_date", "target_date", "horizon", "target_cases", "population",
            "outbreak_threshold", "observed_outbreak", "outbreak_probability",
            "outbreak_decision_threshold", "predicted_mean", "pred_mean_incidence",
            "risk_abs_threshold", "risk_z_threshold", "flag_current_probability",
            "flag_risk_abs", "flag_risk_z", "flag_pred_mean", "methods_count"]
    test[keep].to_csv(args.output / "test_alert_predictions.csv", index=False)
    validation[keep].to_csv(args.output / "validation_alert_predictions.csv", index=False)
    plot_metrics(test_selected, args.output / "outbreak_method_comparison.png")
    write_report(args.title, test_selected, policies, args.output, args.run_mode)
    print(test_selected.groupby("method")[["precision", "specificity", "sensitivity", "f1"]].mean().sort_values("f1", ascending=False).to_string())


if __name__ == "__main__":
    main()
