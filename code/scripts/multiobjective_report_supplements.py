#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Emu, Pt

from dengue_h3.evaluation import _coverage, _crps, outbreak_classification_metrics


METRICS = (
    "mae",
    "rmse",
    "crps",
    "outbreak_brier",
    "outbreak_precision",
    "outbreak_specificity",
    "outbreak_sensitivity",
    "outbreak_f1",
)

LOSS_METRICS = ("mae", "rmse", "crps", "outbreak_brier")
IDENTITY_COLUMNS = ("fcode", "issue_date", "target_date", "horizon")


def add_table(
    document: Document,
    frame: pd.DataFrame,
    digits: int = 3,
    widths: tuple[int, ...] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, frame.columns):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = (
                f"{value:.{digits}f}"
                if isinstance(value, (float, np.floating))
                else str(value)
            )
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Emu(width)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def load_rolling_results(
    directory: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    metrics_path = directory / "metrics.csv"
    comparison_path = directory / "baseline_vs_rolling_mean_H1_H3.csv"
    manifest_path = directory / "run_manifest.json"
    missing = [
        path
        for path in (metrics_path, comparison_path, manifest_path)
        if not path.is_file()
    ]
    if missing:
        names = ", ".join(str(path) for path in missing)
        raise FileNotFoundError(f"Missing rolling-refit output files: {names}")

    metrics = pd.read_csv(metrics_path)
    comparison = pd.read_csv(comparison_path)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    required_metrics = {"horizon", *METRICS}
    missing_metrics = required_metrics.difference(metrics.columns)
    if missing_metrics:
        raise ValueError(
            f"Rolling-refit metrics are missing columns: {sorted(missing_metrics)}"
        )
    horizons = set(metrics["horizon"].astype(int))
    if horizons != {1, 2, 3}:
        raise ValueError(f"Expected rolling-refit horizons 1-3, found {sorted(horizons)}")
    required_comparison = {"approach", *METRICS}
    missing_comparison = required_comparison.difference(comparison.columns)
    if missing_comparison:
        raise ValueError(
            "Rolling-refit comparison is missing columns: "
            f"{sorted(missing_comparison)}"
        )
    approaches = set(comparison["approach"])
    expected_approaches = {
        "Single refit on 2010-2020",
        "Rolling refit at each issue date",
    }
    if approaches != expected_approaches:
        raise ValueError(f"Unexpected rolling-refit approaches: {sorted(approaches)}")
    if not isinstance(manifest.get("vintages_used"), list):
        raise ValueError("Rolling-refit manifest has no vintages_used list")
    if not isinstance(manifest.get("predictive_draws"), int):
        raise ValueError("Rolling-refit manifest has no predictive_draws count")
    return metrics, comparison, manifest


def _predictive_draws(directory: Path, horizon: int) -> np.ndarray:
    """Load one horizon without copying the large INLA binary arrays."""
    npz_path = directory / "predictive_samples.npz"
    if npz_path.is_file():
        with np.load(npz_path) as archive:
            cube = archive["samples"]
        if cube.ndim != 4 or cube.shape[2] != 3:
            raise ValueError(f"Unexpected predictive sample shape in {npz_path}: {cube.shape}")
        return cube[:, :, horizon - 1, :].reshape(-1, cube.shape[-1])

    shape_path = directory / f"samples_H{horizon}_shape.csv"
    binary_path = directory / f"samples_H{horizon}.bin"
    if not shape_path.is_file() or not binary_path.is_file():
        raise FileNotFoundError(
            f"Missing predictive draws for H{horizon} in {directory}"
        )
    shape_row = pd.read_csv(shape_path).iloc[0]
    shape = tuple(int(value) for value in shape_row)
    return np.memmap(
        binary_path,
        dtype=np.float64,
        mode="r",
        shape=shape,
        order="F",
    )


def _score_period_horizon(
    frame: pd.DataFrame,
    draws: np.ndarray,
    decision_threshold: float,
    model: str,
) -> dict[str, object]:
    population = frame["population"].to_numpy(float)
    observed_cases = frame["target_cases"].to_numpy(float)
    observed_incidence = observed_cases / population * 100_000.0
    sample_incidence = draws / population[:, None] * 100_000.0
    error = sample_incidence.mean(axis=1) - observed_incidence
    outbreak_threshold = frame["outbreak_threshold"].to_numpy(float)
    outbreak_probability = frame["outbreak_probability"].to_numpy(float)
    observed_outbreak = observed_cases > outbreak_threshold
    classification = outbreak_classification_metrics(
        observed_cases,
        outbreak_threshold,
        outbreak_probability,
        decision_threshold=decision_threshold,
    )
    return {
        "model": model,
        "split": "test_2023_2025",
        "horizon": int(frame["horizon"].iloc[0]),
        "n": len(frame),
        "mae": float(np.abs(error).mean()),
        "rmse": float(np.sqrt(np.square(error).mean())),
        "bias": float(error.mean()),
        "crps": float(_crps(observed_incidence, sample_incidence).mean()),
        "coverage_50": _coverage(observed_incidence, sample_incidence, 0.50),
        "coverage_80": _coverage(observed_incidence, sample_incidence, 0.80),
        "coverage_95": _coverage(observed_incidence, sample_incidence, 0.95),
        "outbreak_brier": float(
            np.square(outbreak_probability - observed_outbreak).mean()
        ),
        **classification,
    }


def build_period_metrics(
    sources: dict[str, Path],
    output: Path,
    start_date: str = "2023-01-01",
    end_date: str = "2025-12-01",
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Rescore stored forecasts/draws on an inclusive target-date window."""
    start = pd.Timestamp(start_date)
    end = pd.Timestamp(end_date)
    rows: list[dict[str, object]] = []
    for model, directory in sources.items():
        predictions_path = directory / "predictions.csv"
        predictions = pd.read_csv(predictions_path, dtype={"fcode": str})
        predictions["target_date"] = pd.to_datetime(predictions["target_date"])
        missing = set(IDENTITY_COLUMNS).difference(predictions.columns)
        if missing:
            raise ValueError(
                f"Prediction file {predictions_path} is missing {sorted(missing)}"
            )
        for horizon in (1, 2, 3):
            full = predictions[predictions["horizon"].eq(horizon)].reset_index(drop=True)
            draws = _predictive_draws(directory, horizon)
            if len(full) != len(draws):
                raise ValueError(
                    f"Prediction/draw row mismatch for {model} H{horizon}: "
                    f"{len(full)} != {len(draws)}"
                )
            mask = full["target_date"].between(start, end).to_numpy()
            period = full.loc[mask].reset_index(drop=True)
            if period.empty:
                raise ValueError(
                    f"No predictions for {model} H{horizon} in {start_date}--{end_date}"
                )
            if "outbreak_decision_threshold" in period:
                thresholds = period["outbreak_decision_threshold"].dropna().unique()
                if len(thresholds) != 1:
                    raise ValueError(
                        f"Expected one decision threshold for {model} H{horizon}"
                    )
                decision_threshold = float(thresholds[0])
            else:
                decision_threshold = 0.50
            rows.append(
                _score_period_horizon(
                    period,
                    draws[mask],
                    decision_threshold,
                    model,
                )
            )

    by_horizon = pd.DataFrame(rows).sort_values(["model", "horizon"])
    summary = by_horizon.groupby("model", as_index=False)[list(METRICS)].mean()
    for metric in METRICS:
        summary[f"rank_{metric}"] = summary[metric].rank(
            ascending=metric in LOSS_METRICS,
            method="min",
        )
    rank_columns = [f"rank_{metric}" for metric in METRICS]
    summary["mean_metric_rank"] = summary[rank_columns].mean(axis=1)
    summary["worst_metric_rank"] = summary[rank_columns].max(axis=1)
    summary = summary.sort_values(["mean_metric_rank", "worst_metric_rank"])
    by_horizon.to_csv(output / "final_metrics_by_horizon_2023_2025.csv", index=False)
    summary.to_csv(output / "final_metrics_mean_H1_H3_2023_2025.csv", index=False)
    return by_horizon, summary


def add_restricted_test_section(
    document: Document,
    by_horizon: pd.DataFrame,
    summary: pd.DataFrame,
    figure_path: Path,
) -> None:
    """Append the 2023--2025 sensitivity analysis without replacing prior results."""
    _add_page_break(document)
    document.add_heading("10. Sensitivity analysis excluding 2021 and 2022", level=1)
    document.add_paragraph(
        "This supplementary analysis retains the original fitted models, predictive draws, calibration, "
        "and validation-selected alert thresholds, but restricts scoring to target months from January "
        "2023 through December 2025. It therefore removes 2021 and 2022 from the reported test window "
        "without retraining or reselecting any model. Each horizon contains 36 months × 554 communes = "
        "19,944 forecasts. The original 2021–2025 results in Section 5 remain unchanged."
    )

    display = summary[
        [
            "model",
            "mae",
            "rmse",
            "crps",
            "outbreak_brier",
            "outbreak_precision",
            "outbreak_specificity",
            "outbreak_sensitivity",
            "outbreak_f1",
        ]
    ].copy()
    display.columns = [
        "Model",
        "MAE",
        "RMSE",
        "CRPS",
        "Brier",
        "Precision",
        "Specificity",
        "Sensitivity",
        "F1",
    ]
    document.add_heading("10.1 Mean H1–H3 results for 2023–2025", level=2)
    add_table(document, display)

    best_balance = summary.iloc[0]["model"]
    document.add_paragraph(
        f"Using the same equal-weight metric-ranking summary as in the main analysis, {best_balance} "
        "has the best overall rank in this restricted period. These values are a temporal sensitivity "
        "analysis and do not replace the prespecified full-period test result."
    )

    document.add_heading("10.2 H3 results and visualization for 2023–2025", level=2)
    h3_order = [
        "INLA-TS",
        "TCN-GNN Base",
        "Representation-INLA (variant)",
        "DeepST-INLA (proposed hybrid)",
        "DeepST-Alert (proposed deep)",
    ]
    h3 = (
        by_horizon[by_horizon["horizon"].eq(3)]
        .set_index("model")
        .loc[h3_order]
        .reset_index()
    )
    h3_display = h3[
        [
            "model",
            "mae",
            "rmse",
            "crps",
            "outbreak_brier",
            "outbreak_precision",
            "outbreak_specificity",
            "outbreak_sensitivity",
            "outbreak_f1",
        ]
    ].copy()
    h3_display.columns = display.columns
    add_table(document, h3_display)
    document.add_picture(str(figure_path), width=Emu(6_400_800))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph(
        "Figure 3. H3 forecasts restricted to January 2023–December 2025 for the same four "
        "illustrative communes used in the main analysis."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8.5)


def _add_page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_rolling_refit_section(
    document: Document,
    metrics: pd.DataFrame,
    comparison: pd.DataFrame,
    manifest: dict[str, object],
) -> None:
    _add_page_break(document)
    document.add_heading("7. Rolling refit robustness experiment", level=1)
    document.add_paragraph(
        "This supplementary experiment leaves the model-development conclusions unchanged. Architecture, "
        "features, distribution-combination settings, calibration coefficients, and alert thresholds "
        "remain those selected from the development/validation period."
    )
    document.add_paragraph(
        "For the operational rolling-refit evaluation, the deep components were refitted at each issue "
        "date using only observations available by that date. Representation-augmented INLA was rebuilt "
        "from the corresponding causal rows, and the same previously selected decision policy was applied "
        "without reselection. Forecasts were scored for targets from January 2021 through December 2025 "
        "across all 554 communes and H1–H3."
    )
    display = metrics[
        [
            "horizon",
            "mae",
            "rmse",
            "crps",
            "outbreak_brier",
            "outbreak_precision",
            "outbreak_specificity",
            "outbreak_sensitivity",
            "outbreak_f1",
        ]
    ].copy()
    display["horizon"] = display["horizon"].map(lambda value: f"H{int(value)}")
    mean_row = pd.DataFrame(
        [["Mean", *[metrics[column].mean() for column in METRICS]]],
        columns=display.columns,
    )
    display = pd.concat([display, mean_row], ignore_index=True)
    display.columns = [
        "Horizon",
        "MAE",
        "RMSE",
        "CRPS",
        "Brier",
        "Precision",
        "Specificity",
        "Sensitivity",
        "F1",
    ]
    add_table(
        document,
        display,
        widths=(
            662_940,
            731_520,
            731_520,
            731_520,
            731_520,
            731_520,
            731_520,
            731_520,
            607_695,
        ),
    )
    document.add_paragraph()
    indexed = comparison.set_index("approach")
    baseline = indexed.loc["Single refit on 2010-2020"]
    rolling = indexed.loc["Rolling refit at each issue date"]
    deterioration = {
        "MAE": 100.0 * (rolling.mae - baseline.mae) / baseline.mae,
        "RMSE": 100.0 * (rolling.rmse - baseline.rmse) / baseline.rmse,
        "CRPS": 100.0 * (rolling.crps - baseline.crps) / baseline.crps,
        "Brier": 100.0
        * (rolling.outbreak_brier - baseline.outbreak_brier)
        / baseline.outbreak_brier,
        "F1": 100.0
        * (rolling.outbreak_f1 - baseline.outbreak_f1)
        / baseline.outbreak_f1,
    }
    f1_change = f"{deterioration['F1']:+.1f}%".replace("-", "−")
    document.add_paragraph(
        "Relative to the original single-refit test run, the rolling-refit result was slightly worse "
        f"overall: MAE +{deterioration['MAE']:.1f}%, RMSE +{deterioration['RMSE']:.1f}%, "
        f"CRPS +{deterioration['CRPS']:.1f}%, Brier +{deterioration['Brier']:.1f}%, and F1 "
        f"{f1_change} when lower-is-better metrics are expressed as deterioration. The experiment used "
        f"{len(manifest['vintages_used'])} issue-date vintages and {manifest['predictive_draws']} "
        "predictive draws per forecast. It is a robustness check, not a second model-selection stage."
    )


def _add_box(axis, xy, width, height, title, lines, color) -> None:
    patch = FancyBboxPatch(
        xy,
        width,
        height,
        boxstyle="round,pad=0.015,rounding_size=0.02",
        linewidth=1.4,
        edgecolor=color,
        facecolor=color + "18",
    )
    axis.add_patch(patch)
    x, y = xy
    axis.text(
        x + width / 2,
        y + height * 0.70,
        title,
        ha="center",
        va="center",
        fontsize=9.5,
        fontweight="bold",
        color=color,
    )
    axis.text(
        x + width / 2,
        y + height * 0.35,
        "\n".join(lines),
        ha="center",
        va="center",
        fontsize=7.5,
        color="#263238",
    )


def _add_arrow(axis, start, end) -> None:
    axis.add_patch(
        FancyArrowPatch(
            start,
            end,
            arrowstyle="-|>",
            mutation_scale=12,
            linewidth=1.2,
            color="#607D8B",
        )
    )


def build_architecture_figure(path: Path) -> None:
    figure, axis = plt.subplots(figsize=(15, 9.7))
    axis.set_xlim(0, 1)
    axis.set_ylim(0, 1)
    axis.axis("off")
    axis.text(
        0.5,
        0.975,
        "DeepST-INLA: neural architecture and probabilistic outbreak decision",
        ha="center",
        va="center",
        fontsize=16,
        fontweight="bold",
    )
    axis.text(0.01, 0.91, "A", fontsize=15, fontweight="bold")
    axis.text(
        0.045,
        0.91,
        "Shared M5 TCN-GNN architecture",
        fontsize=11,
        fontweight="bold",
    )
    _add_box(
        axis,
        (0.02, 0.72),
        0.10,
        0.13,
        "Monthly inputs",
        ["12 months", "5 local + 7 climate", "12 channels"],
        "#1565C0",
    )
    _add_box(
        axis,
        (0.15, 0.72),
        0.09,
        0.13,
        "Input projection",
        ["Conv1D 1×1", "12 → 32"],
        "#C49000",
    )
    for index, dilation in enumerate((1, 2, 4)):
        _add_box(
            axis,
            (0.27 + 0.105 * index, 0.72),
            0.085,
            0.13,
            f"TCN block {index + 1}",
            ["Conv1D k=3", "32 → 32", f"dilation={dilation}", "residual"],
            "#C49000",
        )
    _add_box(
        axis,
        (0.60, 0.80),
        0.09,
        0.09,
        "Commune ID",
        ["embedding 554×8"],
        "#455A64",
    )
    _add_box(
        axis,
        (0.60, 0.68),
        0.09,
        0.09,
        "Conditioning",
        ["Linear 8 → 64", "scale + shift"],
        "#455A64",
    )
    _add_box(
        axis,
        (0.72, 0.72),
        0.10,
        0.13,
        "Graph attention",
        ["1 layer", "4 heads × 8", "→ 32"],
        "#1565C0",
    )
    _add_box(
        axis,
        (0.85, 0.72),
        0.11,
        0.13,
        "Residual spatial state",
        ["local + spatial", "32 dimensions"],
        "#2E7D32",
    )
    _add_box(
        axis,
        (0.04, 0.56),
        0.12,
        0.10,
        "Seasonal memory",
        ["14 → 32", "ReLU + dropout"],
        "#6A1B9A",
    )
    _add_box(
        axis,
        (0.19, 0.56),
        0.10,
        0.10,
        "Horizon ID",
        ["embedding 3×8"],
        "#6A1B9A",
    )
    _add_box(
        axis,
        (0.33, 0.56),
        0.13,
        0.10,
        "Concatenate",
        ["32 spatial + 32 seasonal", "+ 8 horizon = 72"],
        "#AD1457",
    )
    _add_box(
        axis,
        (0.50, 0.56),
        0.10,
        0.10,
        "Dense",
        ["72 → 32"],
        "#C49000",
    )
    _add_box(
        axis,
        (0.64, 0.54),
        0.20,
        0.13,
        "Horizon-specific heads",
        ["log-location 32 → 1", "scale 32 → 1", "outbreak logit 32 → 1", "H1, H2, H3"],
        "#1565C0",
    )
    for start, end in (
        ((0.12, 0.785), (0.15, 0.785)),
        ((0.24, 0.785), (0.27, 0.785)),
        ((0.355, 0.785), (0.375, 0.785)),
        ((0.46, 0.785), (0.48, 0.785)),
        ((0.565, 0.785), (0.60, 0.785)),
        ((0.69, 0.745), (0.72, 0.785)),
        ((0.82, 0.785), (0.85, 0.785)),
        ((0.16, 0.61), (0.33, 0.61)),
        ((0.29, 0.61), (0.33, 0.61)),
        ((0.46, 0.61), (0.50, 0.61)),
        ((0.60, 0.61), (0.64, 0.61)),
    ):
        _add_arrow(axis, start, end)
    axis.text(
        0.5,
        0.515,
        "Base and focal-warning experts share this 22,269-parameter architecture but are trained independently.",
        ha="center",
        va="center",
        fontsize=8,
        style="italic",
    )
    axis.text(0.01, 0.47, "B", fontsize=15, fontweight="bold")
    axis.text(
        0.045,
        0.47,
        "DeepST-INLA predictive distribution",
        fontsize=11,
        fontweight="bold",
    )
    distribution_boxes = (
        (0.02, "Base representation", ["32-dimensional state"], "#1565C0"),
        (0.15, "Prepare for INLA", ["standardize + residualize", "against 21 temporal terms", "PCA 32 → 4"], "#6A1B9A"),
        (0.31, "Representation-INLA", ["4 deep PCs", "temporal + BYM2", "negative binomial"], "#C49000"),
        (0.48, "INLA posterior", ["500 predictive draws"], "#1565C0"),
        (0.64, "Risk-gated mixture", ["validation-selected expert", "base or warning draws", "partial replacement"], "#C62828"),
        (0.82, "Final distribution", ["500 predictive draws"], "#1565C0"),
    )
    for x, title, lines, color in distribution_boxes:
        _add_box(axis, (x, 0.34), 0.13, 0.10, title, lines, color)
    for left, right in ((0.15, 0.15), (0.28, 0.31), (0.44, 0.48), (0.61, 0.64), (0.77, 0.82)):
        _add_arrow(axis, (left, 0.39), (right, 0.39))
    _add_box(
        axis,
        (0.15, 0.22),
        0.13,
        0.08,
        "Base count expert",
        ["count NLL only", "500 draws"],
        "#2E7D32",
    )
    _add_box(
        axis,
        (0.32, 0.22),
        0.13,
        0.08,
        "Focal-warning expert",
        ["count NLL + focal BCE", "500 draws + p_aux"],
        "#2E7D32",
    )
    _add_arrow(axis, (0.28, 0.26), (0.64, 0.37))
    _add_arrow(axis, (0.45, 0.26), (0.64, 0.37))
    axis.text(0.01, 0.16, "C", fontsize=15, fontweight="bold")
    axis.text(
        0.045,
        0.16,
        "Outbreak probability and alert",
        fontsize=11,
        fontweight="bold",
    )
    alert_boxes = (
        (0.04, "Threshold T(c,m)", ["5-year same-month", "mean + 2 × sample SD"], "#455A64"),
        (0.22, "p_count", ["fraction of final draws", "above T(c,m)"], "#1565C0"),
        (0.40, "Validation alert score", ["p_raw = w_count p_count", "+ w_aux p_aux"], "#C49000"),
        (0.59, "Platt calibration", ["sigmoid(a + b logit(p_raw))"], "#1565C0"),
        (0.77, "Alert cutoff", ["H1 .33; H2 .35; H3 .37"], "#455A64"),
        (0.91, "Outbreak alert", ["0 / 1"], "#2E7D32"),
    )
    widths = (0.14, 0.14, 0.15, 0.15, 0.11, 0.08)
    for (x, title, lines, color), width in zip(alert_boxes, widths):
        _add_box(axis, (x, 0.04), width, 0.09, title, lines, color)
    for left, right in ((0.18, 0.22), (0.36, 0.40), (0.55, 0.59), (0.74, 0.77), (0.88, 0.91)):
        _add_arrow(axis, (left, 0.085), (right, 0.085))
    _add_arrow(axis, (0.885, 0.34), (0.66, 0.13))
    _add_arrow(axis, (0.385, 0.22), (0.475, 0.13))
    figure.savefig(path, dpi=250, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def add_architecture_and_outbreak_sections(
    document: Document,
    figure_path: Path,
    selection_path: Path,
) -> None:
    selection = json.loads(selection_path.read_text(encoding="utf-8"))
    policies = selection["policies"]
    document.add_heading("8. Detailed architecture of the current model", level=1)
    document.add_picture(
        str(figure_path),
        width=Emu(6_492_240),
        height=Emu(4_194_176),
    )
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph(
        "Figure 2. Block architecture of the shared M5 neural expert, Representation-INLA distribution, "
        "validation-selected neural expert and risk-gated outbreak decision path."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8.5)
    inputs = document.add_paragraph()
    inputs.add_run(
        "Each neural expert receives a 12-month sequence. At every month, five local inputs—log cases, "
        "log incidence, log population, and sine/cosine month encodings—are concatenated with seven "
        "climate variables: temperature, relative humidity, rainfall, soil "
    )
    inputs.add_run(
        "moisture, solar radiation, dew point, and ET0. This produces 12 channels per month, projected "
        "to 32 hidden units by a 1×1 convolution."
    )
    document.add_paragraph(
        "The temporal encoder has three residual causal convolution blocks. Each block uses a 32-to-32 "
        "Conv1D with kernel size 3, GroupNorm, ReLU and dropout 0.10; dilations are 1, 2 and 4. The graph "
        "stage is a single multi-head attention layer with four 8-dimensional heads. Its output is added "
        "residually to the local state. The selected M5 base does not use the optional dynamic spatial gate."
    )
    document.add_paragraph(
        "Commune conditioning uses a learned 8-dimensional embedding for each of 554 communes. A linear "
        "8-to-64 transform provides scale and shift vectors for the 32-unit local state. The target-aligned "
        "seasonal branch has 14 inputs: five log-transformed same-month case counts from 1–5 years earlier, "
        "five availability indicators, plus their mean, sample standard deviation, trend and zero-share. "
        "For each horizon, the 32-unit spatial state, 32-unit seasonal state and 8-unit horizon embedding "
        "are concatenated (72 units) and mapped to a 32-dimensional representation."
    )
    document.add_paragraph(
        "Each horizon has three 32-to-1 heads: log-location, log-scale and outbreak logit. For the "
        "lognormal count model, softplus constrains the scale to be positive, and 500 predictive case-count "
        "draws are sampled. The focal-warning expert has exactly the same 22,269-parameter architecture but "
        "is trained independently with the additional focal-BCE objective. Its count draws remain a "
        "candidate distribution, while its auxiliary outbreak probability p_aux is also used as the risk "
        "signal and as one member of the final alert score."
    )
    architecture = pd.DataFrame(
        [
            ("Input projection", "Conv1D, kernel 1", "12 → 32", "416"),
            (
                "Causal temporal encoder",
                "3 residual Conv1D blocks",
                "32; k=3; d=1,2,4",
                "9,504",
            ),
            ("Seasonal encoder", "Linear + ReLU + dropout", "14 → 32", "480"),
            (
                "Commune conditioning",
                "554 embeddings + scale/shift",
                "8; then 8 → 64",
                "5,008",
            ),
            ("Graph attention", "1 layer, 4 heads", "4 × 8 → 32", "4,204"),
            (
                "Horizon representation",
                "Embedding + dense layer",
                "8; then 72 → 32",
                "2,360",
            ),
            ("Output heads", "3 heads × 3 horizons", "32 → 1 each", "297"),
            ("One neural expert", "M5 TCN-GNN", "hidden width 32", "22,269"),
        ],
        columns=["Component", "Layers", "Dimensions/settings", "Parameters"],
    )
    add_table(
        document,
        architecture,
        widths=(1_371_600, 1_554_480, 1_737_360, 731_520),
    )
    document.add_paragraph()
    document.add_paragraph(
        "For DeepST-INLA, the 32-dimensional deep representation is standardized, residualized against "
        "21 temporal covariates already present in the INLA design, and reduced to four principal "
        "components. Those four residual deep components enter the negative-binomial INLA model together "
        "with temporal effects and a BYM2 spatial effect. This separation is deliberate: the neural "
        "network learns nonlinear residual information, while INLA contributes structured temporal/spatial "
        "regularisation and a posterior predictive count distribution."
    )

    document.add_heading("9. Outbreak probability and alert calculation", level=1)
    document.add_paragraph(
        "The alert calculation is designed to answer one practical question: how likely is the predicted "
        "case count to represent an outbreak? The procedure first defines the outbreak threshold, then "
        "estimates risk in two complementary ways, combines the two probabilities, and finally calibrates "
        "the combined value. All combination weights, calibration coefficients, and decision thresholds "
        "are selected using validation data only."
    )
    threshold_heading = document.add_heading(
        "9.1 Define the outbreak threshold",
        level=2,
    )
    threshold_heading.runs[0].bold = False
    threshold_explanation = document.add_paragraph(
        "A single fixed case threshold is not appropriate because normal dengue incidence differs across "
        "communes and seasons. For commune c and target month m, the threshold is therefore calculated "
        "from the same calendar month in the previous five years:"
    )
    threshold_explanation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula = document.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula.add_run(
        "T(c,m) = mean(y₁,…,y₅) + 2 × sample SD(y₁,…,y₅)"
    ).bold = True
    outbreak_label = document.add_paragraph(
        "At least two historical observations are required. The observed target month is labelled as an "
        "outbreak only when its observed case count is strictly greater than T(c,m)."
    )
    outbreak_label.alignment = WD_ALIGN_PARAGRAPH.CENTER

    probability_heading = document.add_heading(
        "9.2 Estimate outbreak probability from two sources",
        level=2,
    )
    probability_heading.runs[0].bold = False
    document.add_paragraph(
        "First, the model uses its predictive case-count distribution. With 500 predictive draws, p_count "
        "is the fraction of draws that exceed the outbreak threshold: p_count = number of draws above "
        "T(c,m) / 500. This uses the uncertainty of the full forecast rather than relying only on the "
        "predicted mean. For DeepST-INLA these draws come from the final INLA-supported distribution; for "
        "DeepST-Alert they come directly from the deep model."
    )
    document.add_paragraph(
        "Second, p_aux is the probability that the deep model predicts directly for the outbreak class. "
        "This additional signal is useful because a model can forecast average case counts reasonably well "
        "yet still miss difficult or relatively rare outbreaks. Focal BCE is used during training so these "
        "outbreak cases receive greater attention."
    )
    document.add_paragraph(
        "The two signals are combined as p_raw = w_count × p_count + w_aux × p_aux. Combining probabilities "
        "does not guarantee that the result is well calibrated—for example, p_raw = 0.70 does not "
        "automatically mean that about 70% of comparable forecasts become outbreaks. Platt scaling is "
        "therefore fitted on validation data: p_final = sigmoid(a + b × logit(p_raw)). The calibrated "
        "probability is then compared with a horizon-specific alert threshold. The selected validation "
        "parameters are shown below."
    )
    policy_rows = []
    for horizon in ("1", "2", "3"):
        alert = policies[horizon]["alert"]
        policy_rows.append(
            (
                f"H{horizon}",
                f"{alert['weights'][0]:.2f}",
                f"{alert['weights'][3]:.2f}",
                f"{alert['platt_intercept']:.4f}",
                f"{alert['platt_slope']:.4f}",
                f"{alert['decision_threshold']:.2f}",
            )
        )
    policy = pd.DataFrame(
        policy_rows,
        columns=[
            "Horizon",
            "p_count weight",
            "p_aux weight",
            "Platt intercept a",
            "Platt slope b",
            "Alert threshold",
        ],
    )
    add_table(
        document,
        policy,
        widths=(640_080, 1_051_560, 960_120, 1_051_560, 914_400, 914_400),
    )
    document.add_paragraph()
    document.add_paragraph(
        "The H1–H3 weights, Platt coefficients, and alert thresholds above were selected from the "
        "2019–2020 validation period and then fixed before evaluation on the 2021–2025 test period."
    )

    example_heading = document.add_heading("9.3 Worked example (H2)", level=2)
    example_heading.runs[0].bold = False
    document.add_paragraph(
        "Assume one commune recorded 10, 14, 12, 16, and 18 cases in the same calendar month during the "
        "previous five years. The mean is 14 and the sample standard deviation is 3.162, giving "
        "T(c,m) = 14 + 2 × 3.162 = 20.325 cases."
    )
    document.add_paragraph(
        "Suppose 350 of the 500 H2 predictive draws exceed 20.325. Then p_count = 350/500 = 0.70. "
        "Separately, the deep model directly predicts p_aux = 0.80. For H2, validation selected weights "
        "of 0.40 for p_count and 0.60 for p_aux, so p_raw = 0.40 × 0.70 + 0.60 × 0.80 = 0.76."
    )
    document.add_paragraph(
        "Because p_raw is not assumed to be perfectly calibrated, the H2 Platt transformation is applied: "
        "p_final = sigmoid(−0.7003 + 2.3340 × logit(0.76)) = 0.880. The H2 alert threshold is 0.35, "
        "so 0.880 ≥ 0.35 and the model issues an outbreak alert. If the observed target count were 27, "
        "the observed outbreak label would also be 1 because 27 > 20.325."
    )
    reading_guide = document.add_paragraph()
    reading_guide.add_run("Reading guide. ").bold = True
    reading_guide.add_run(
        "Threshold first; then p_count from the predictive count distribution and p_aux from the direct "
        "outbreak prediction; then combine the two; then use Platt scaling to calibrate the combined "
        "probability; finally compare it with the validation-selected alert threshold."
    )
