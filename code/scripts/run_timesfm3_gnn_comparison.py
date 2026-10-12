#!/usr/bin/env python3
"""Evaluate official TimesFM 3 and a frozen-TimesFM graph adapter.

The experiment keeps the report's original split and uses TimesFM 3 with the
seven selected climate variables as past-only covariates.  TimesFM-GNN freezes
the foundation model and trains a small graph-attention adapter on 2011-2018,
selects its epoch/alert cutoffs on 2019-2020, refits through 2020, and evaluates
the locked 2021-2025 targets.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import time
from copy import deepcopy
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt
from torch import nn

from dengue_h3.distributions import predictive_sample
from dengue_h3.data import _moh_outbreak_threshold_lookup
from dengue_h3.evaluation import outbreak_classification_metrics, score_predictions
from dengue_h3.graph import build_candidate_graph
from dengue_h3.models import MultiHeadGraphLayer
from multiobjective_report_supplements import add_table, build_period_metrics


ROOT = Path(__file__).resolve().parents[2]
DATA_PATH = ROOT / "code/data/processed/feature_suite/monthly_commune_open.csv"
CENTROIDS_PATH = ROOT / "code/data/processed/qc_v1/commune_centroids.csv"
ADJACENCY_PATH = ROOT / "code/data/processed/qc_v1/administrative_adjacency.csv"
MOH_THRESHOLDS_PATH = ROOT / "code/data/raw/operational_2026/commune_monthly_thresholds.csv"
CURRENT_REPORT = ROOT / "results/current"
RESULTS = ROOT / "results/timesfm3_comparison_2021_2025"
COPY_REPORT = ROOT / "results/timesfm3_report"
CHECKPOINT = "google/timesfm-3.0-pytorch"
CLIMATE_COLUMNS = (
    "t2m_avg",
    "rh_avg",
    "tp_accum",
    "soil_moisture",
    "solar_downward",
    "dewpoint_c",
    "et0_mm",
)
QUANTILE_LEVELS = np.arange(0.1, 1.0, 0.1, dtype=np.float32)
HORIZONS = (1, 2, 3)
TEST_START = pd.Timestamp("2021-01-01")
TEST_END = pd.Timestamp("2025-12-01")
VALIDATION_START = pd.Timestamp("2019-01-01")
VALIDATION_END = pd.Timestamp("2020-12-01")
TRAIN_END = pd.Timestamp("2018-12-01")
MODEL_LABELS = {
    "TIMESFM-3": "TimesFM 3 (no GNN)",
    "TIMESFM-3-GNN": "TimesFM 3–GNN",
}


def load_panel() -> tuple[pd.DataFrame, tuple[str, ...], pd.DatetimeIndex, np.ndarray, np.ndarray, np.ndarray]:
    frame = pd.read_csv(DATA_PATH, dtype={"fcode": str})
    frame["date"] = pd.to_datetime(
        {"year": frame["year"], "month": frame["month"], "day": 1}
    )
    units = tuple(sorted(frame["fcode"].unique()))
    dates = pd.DatetimeIndex(sorted(frame["date"].unique()))
    if len(units) != 554:
        raise ValueError(f"Expected 554 communes, found {len(units)}")
    cases = (
        frame.pivot(index="date", columns="fcode", values="obs_dengue_cases")
        .loc[dates, list(units)]
        .to_numpy(np.float32)
    )
    population = (
        frame.pivot(index="date", columns="fcode", values="population")
        .loc[dates, list(units)]
        .to_numpy(np.float32)
    )
    climate = np.stack(
        [
            frame.pivot(index="date", columns="fcode", values=column)
            .loc[dates, list(units)]
            .to_numpy(np.float32)
            for column in CLIMATE_COLUMNS
        ],
        axis=-1,
    )
    return frame, units, dates, cases, population, climate


def build_backbone_cache(
    cache_path: Path,
    units: tuple[str, ...],
    dates: pd.DatetimeIndex,
    cases: np.ndarray,
    climate: np.ndarray,
    batch_size: int,
    device: str,
) -> tuple[np.ndarray, np.ndarray]:
    point = np.full((len(dates), len(units), 3), np.nan, dtype=np.float32)
    quantiles = np.full(
        (len(dates), len(units), 3, len(QUANTILE_LEVELS)),
        np.nan,
        dtype=np.float32,
    )
    completed_through = -1
    if cache_path.is_file():
        with np.load(cache_path) as archive:
            cached_dates = pd.to_datetime(archive["issue_dates"])
            if not np.array_equal(cached_dates.to_numpy(), dates.to_numpy()):
                raise ValueError("TimesFM cache dates do not match the current panel")
            point = archive["point"]
            quantiles = archive["quantiles"]
            completed_through = int(archive.get("completed_through", -1))
            if bool(archive.get("complete", False)):
                return point, quantiles

    from timesfm3 import ModelConfig, TimesFM3Evaluator

    model = TimesFM3Evaluator(
        ModelConfig(
            checkpoint_path=CHECKPOINT,
            per_core_batch_size=batch_size,
            device=device,
        )
    )
    # October 2010 is the first issue date needed for a 2011 H3 target.
    first_issue = int(np.flatnonzero(dates >= pd.Timestamp("2010-10-01"))[0])
    # November 2025 is the last origin needed for a target in the locked test.
    last_issue = int(np.flatnonzero(dates <= pd.Timestamp("2025-11-01"))[-1])
    first_issue = max(first_issue, completed_through + 1)
    started = time.time()
    for issue_index in range(first_issue, last_issue + 1):
        contexts = [cases[: issue_index + 1, node][None] for node in range(len(units))]
        past_covariates = [
            climate[: issue_index + 1, node].T for node in range(len(units))
        ]
        outputs = list(
            model.predict_batch(
                contexts,
                horizon=3,
                past_only_covariates=past_covariates,
                return_quantiles=True,
                use_symmetric_averaging=False,
            )
        )
        remaining = min(3, len(dates) - issue_index - 1)
        for node, output in enumerate(outputs):
            point[issue_index, node, :remaining] = np.asarray(output.forecast)[
                0, :remaining
            ]
            quantiles[issue_index, node, :remaining] = np.asarray(output.quantiles)[
                0, :remaining
            ]
        point[issue_index] = np.clip(point[issue_index], 0.0, None)
        quantiles[issue_index] = np.maximum.accumulate(
            np.clip(quantiles[issue_index], 0.0, None), axis=-1
        )
        completed = issue_index - first_issue + 1
        if completed == 1 or completed % 6 == 0 or issue_index == last_issue:
            print(
                f"TimesFM backbone: {dates[issue_index].date()} "
                f"({completed}/{last_issue - first_issue + 1}, "
                f"{time.time() - started:.1f}s)",
                flush=True,
            )
            save_backbone_cache(
                cache_path,
                dates,
                point,
                quantiles,
                completed_through=issue_index,
                complete=issue_index == last_issue,
            )
    del model
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
    return point, quantiles


def save_backbone_cache(
    cache_path: Path,
    dates: pd.DatetimeIndex,
    point: np.ndarray,
    quantiles: np.ndarray,
    *,
    completed_through: int,
    complete: bool,
) -> None:
    """Atomically checkpoint a long CPU inference run for safe resumption."""
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = cache_path.with_name(cache_path.name + ".tmp.npz")
    np.savez_compressed(
        temporary,
        issue_dates=dates.to_numpy(dtype="datetime64[ns]"),
        point=point,
        quantiles=quantiles,
        quantile_levels=QUANTILE_LEVELS,
        completed_through=np.int64(completed_through),
        complete=np.bool_(complete),
    )
    temporary.replace(cache_path)


def moh_thresholds(
    units: tuple[str, ...], dates: pd.DatetimeIndex, cases: np.ndarray
) -> np.ndarray:
    """Return row-level MOH thresholds with an early-history fallback.

    The audited export is authoritative wherever it contains a value.  Before
    that export begins, use the same MOH implementation as the main pipeline;
    rows with fewer than two prior same-month observations remain unavailable
    and are masked from the outbreak loss and metrics.
    """
    panel = pd.DataFrame(
        {
            "date": np.repeat(dates.to_numpy(), len(units)),
            "fcode": np.tile(np.asarray(units), len(dates)),
            "cases": cases.reshape(-1),
        }
    )
    lookup = _moh_outbreak_threshold_lookup(panel, dates, units)
    matrix = np.asarray(
        [[lookup.get((unit, date), np.nan) for unit in units] for date in dates],
        dtype=np.float32,
    )
    frame = pd.read_csv(MOH_THRESHOLDS_PATH, dtype={"fcode": str})
    frame["target_date"] = pd.to_datetime(frame["target_date"])
    audited = (
        frame.pivot(index="target_date", columns="fcode", values="epi_threshold")
        .reindex(index=dates, columns=list(units))
        .to_numpy(np.float32)
    )
    matrix[np.isfinite(audited)] = audited[np.isfinite(audited)]
    return matrix


def quantiles_to_draws(values: np.ndarray, draws: int = 500) -> np.ndarray:
    """Turn TimesFM deciles into a reproducible empirical predictive sample."""
    flat = values.reshape(-1, values.shape[-1])
    output = np.empty((len(flat), draws), dtype=np.float32)
    probabilities = (np.arange(draws, dtype=np.float32) + 0.5) / draws
    for index, row in enumerate(flat):
        lower = max(0.0, float(row[0] - (row[1] - row[0])))
        upper = float(row[-1] + (row[-1] - row[-2]))
        knots = np.concatenate(([lower], row, [upper])).astype(np.float32)
        output[index] = np.interp(
            probabilities,
            np.concatenate(([0.0], QUANTILE_LEVELS, [1.0])),
            knots,
        )
    return output.reshape(*values.shape[:-1], draws)


def target_arrays(
    dates: pd.DatetimeIndex,
    cases: np.ndarray,
    population: np.ndarray,
    thresholds: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    targets = np.full((len(dates), len(cases[0]), 3), np.nan, dtype=np.float32)
    target_population = np.full_like(targets, np.nan)
    target_thresholds = np.full_like(targets, np.nan)
    target_dates = np.full((len(dates), 3), np.datetime64("NaT"), dtype="datetime64[ns]")
    for issue_index in range(len(dates)):
        for horizon_index, horizon in enumerate(HORIZONS):
            target_index = issue_index + horizon
            if target_index >= len(dates):
                continue
            targets[issue_index, :, horizon_index] = cases[target_index]
            target_population[issue_index, :, horizon_index] = population[target_index]
            target_thresholds[issue_index, :, horizon_index] = thresholds[target_index]
            target_dates[issue_index, horizon_index] = dates[target_index]
    return targets, target_population, target_thresholds, target_dates


def backbone_features(
    point: np.ndarray,
    quantiles: np.ndarray,
    cases: np.ndarray,
    population: np.ndarray,
    dates: pd.DatetimeIndex,
) -> np.ndarray:
    features = np.full(
        (*point.shape, 14),
        np.nan,
        dtype=np.float32,
    )
    for issue_index, issue_date in enumerate(dates):
        for horizon_index, horizon in enumerate(HORIZONS):
            target_date = issue_date + pd.DateOffset(months=horizon)
            angle = 2.0 * np.pi * target_date.month / 12.0
            pieces = [
                np.log1p(point[issue_index, :, horizon_index])[:, None],
                np.log1p(quantiles[issue_index, :, horizon_index]),
                np.log1p(cases[issue_index])[:, None],
                np.log1p(population[issue_index])[:, None],
                np.full((len(cases[0]), 1), np.sin(angle), dtype=np.float32),
                np.full((len(cases[0]), 1), np.cos(angle), dtype=np.float32),
            ]
            features[issue_index, :, horizon_index] = np.concatenate(pieces, axis=1)
    return features


class TimesFMGraphAdapter(nn.Module):
    def __init__(self, feature_count: int, graph) -> None:
        super().__init__()
        hidden = 32
        self.input_projection = nn.Sequential(
            nn.Linear(feature_count, hidden), nn.ReLU(), nn.Dropout(0.10)
        )
        self.commune_embedding = nn.Embedding(len(graph.units), 8)
        self.condition = nn.Linear(8, hidden * 2)
        self.graph = MultiHeadGraphLayer(
            hidden,
            torch.from_numpy(graph.mask),
            torch.from_numpy(graph.edge_features),
            0.10,
            4,
        )
        self.horizon_embedding = nn.Embedding(3, 8)
        self.representation = nn.Sequential(
            nn.Linear(hidden + 8, hidden), nn.ReLU(), nn.Dropout(0.10)
        )
        self.location_heads = nn.ModuleList([nn.Linear(hidden, 1) for _ in HORIZONS])
        self.scale_heads = nn.ModuleList([nn.Linear(hidden, 1) for _ in HORIZONS])
        self.outbreak_heads = nn.ModuleList([nn.Linear(hidden, 1) for _ in HORIZONS])

    def forward(self, features: torch.Tensor) -> tuple[torch.Tensor, torch.Tensor, torch.Tensor]:
        batch, nodes, horizons, _ = features.shape
        locations, scales, logits = [], [], []
        node_ids = torch.arange(nodes, device=features.device)
        scale_shift = self.condition(self.commune_embedding(node_ids))
        node_scale, node_shift = scale_shift.chunk(2, dim=-1)
        for horizon_index in range(horizons):
            local = self.input_projection(features[:, :, horizon_index])
            local = (1.0 + 0.1 * torch.tanh(node_scale))[None] * local + node_shift[None]
            spatial = local + self.graph(local)
            horizon = self.horizon_embedding(
                torch.full((nodes,), horizon_index, dtype=torch.long, device=features.device)
            )[None].expand(batch, -1, -1)
            hidden = self.representation(torch.cat([spatial, horizon], dim=-1))
            locations.append(self.location_heads[horizon_index](hidden))
            scales.append(
                (torch.nn.functional.softplus(self.scale_heads[horizon_index](hidden)) + 0.05)
                .clamp_max(3.0)
            )
            logits.append(self.outbreak_heads[horizon_index](hidden))
        return (
            torch.cat(locations, dim=-1),
            torch.cat(scales, dim=-1),
            torch.cat(logits, dim=-1),
        )


def split_element_mask(target_dates: np.ndarray, start: pd.Timestamp, end: pd.Timestamp) -> np.ndarray:
    dates = pd.to_datetime(target_dates.reshape(-1)).to_numpy().reshape(target_dates.shape)
    return (dates >= np.datetime64(start)) & (dates <= np.datetime64(end))


def fit_scale(features: np.ndarray, element_mask: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    selected = features[element_mask[:, None, :].repeat(features.shape[1], axis=1)]
    mean = np.nanmean(selected, axis=0, keepdims=True).astype(np.float32)
    sd = np.nanstd(selected, axis=0, keepdims=True).astype(np.float32)
    sd[sd < 1e-6] = 1.0
    return mean, sd


def masked_loss(
    output: tuple[torch.Tensor, torch.Tensor, torch.Tensor],
    targets: torch.Tensor,
    thresholds: torch.Tensor,
    mask: torch.Tensor,
    positive_weight: torch.Tensor,
) -> torch.Tensor:
    location, scale, logits = output
    transformed = torch.log1p(targets.clamp_min(0))
    count = torch.log(scale) + 0.5 * ((transformed - location) / scale) ** 2
    expanded_mask = mask[:, None, :].expand_as(count) & torch.isfinite(count)
    count_loss = count[expanded_mask].mean()
    finite_alert = expanded_mask & torch.isfinite(thresholds)
    if not bool(finite_alert.any()):
        return count_loss
    labels = (targets > thresholds).to(targets.dtype)
    bce = torch.nn.functional.binary_cross_entropy_with_logits(logits, labels, reduction="none")
    bce = bce * torch.where(labels > 0, positive_weight[None, None], 1.0)
    probability = torch.sigmoid(logits)
    correct = torch.where(labels > 0, probability, 1.0 - probability)
    focal = bce * (1.0 - correct).pow(2.0)
    return count_loss + 0.25 * focal[finite_alert].mean()


def fit_adapter(
    features: np.ndarray,
    targets: np.ndarray,
    thresholds: np.ndarray,
    train_mask: np.ndarray,
    validation_mask: np.ndarray,
    graph,
    *,
    fixed_epochs: int | None = None,
    device_name: str = "cpu",
) -> tuple[TimesFMGraphAdapter, list[dict[str, float]], tuple[np.ndarray, np.ndarray], int]:
    seed = 9411
    torch.manual_seed(seed)
    np.random.seed(seed)
    mean, sd = fit_scale(features, train_mask)
    scaled = np.nan_to_num((features - mean) / sd).astype(np.float32)
    train_labels = targets > thresholds
    positive = ((train_labels) & train_mask[:, None, :]).sum(axis=(0, 1))
    negative = ((~train_labels) & train_mask[:, None, :] & np.isfinite(thresholds)).sum(axis=(0, 1))
    positive_weight = np.clip(negative / np.maximum(positive, 1), 1.0, 8.0).astype(np.float32)
    device = torch.device(device_name)
    model = TimesFMGraphAdapter(features.shape[-1], graph).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=1e-3, weight_decay=1e-4)
    tensors = (
        torch.from_numpy(scaled),
        torch.from_numpy(targets),
        torch.from_numpy(thresholds),
        torch.from_numpy(positive_weight).to(device),
    )
    indices = np.flatnonzero(train_mask.any(axis=1))
    validation_indices = np.flatnonzero(validation_mask.any(axis=1))
    rng = np.random.default_rng(seed)
    best_state = deepcopy(model.state_dict())
    best_loss = float("inf")
    best_epoch = 1
    stale = 0
    history: list[dict[str, float]] = []
    epochs = fixed_epochs or 35
    for epoch in range(epochs):
        model.train()
        batch_losses = []
        shuffled = rng.permutation(indices)
        for start in range(0, len(indices), 8):
            batch = shuffled[start : start + 8]
            index = torch.from_numpy(batch)
            optimizer.zero_grad(set_to_none=True)
            loss = masked_loss(
                model(tensors[0][index].to(device)),
                tensors[1][index].to(device),
                tensors[2][index].to(device),
                torch.from_numpy(train_mask[batch]).to(device),
                tensors[3],
            )
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), 1.0)
            optimizer.step()
            batch_losses.append(float(loss.detach().cpu()))
        validation_loss = float("nan")
        if fixed_epochs is None:
            model.eval()
            values = []
            with torch.no_grad():
                for start in range(0, len(validation_indices), 8):
                    batch = validation_indices[start : start + 8]
                    index = torch.from_numpy(batch)
                    values.append(
                        float(
                            masked_loss(
                                model(tensors[0][index].to(device)),
                                tensors[1][index].to(device),
                                tensors[2][index].to(device),
                                torch.from_numpy(validation_mask[batch]).to(device),
                                tensors[3],
                            ).cpu()
                        )
                    )
            validation_loss = float(np.mean(values))
            if validation_loss < best_loss - 1e-4:
                best_loss = validation_loss
                best_state = deepcopy(model.state_dict())
                best_epoch = epoch + 1
                stale = 0
            else:
                stale += 1
        else:
            best_state = deepcopy(model.state_dict())
            best_epoch = epoch + 1
        history.append(
            {
                "epoch": epoch + 1,
                "training_loss": float(np.mean(batch_losses)),
                "validation_loss": validation_loss,
            }
        )
        print(
            f"TimesFM-GNN epoch {epoch + 1}/{epochs}: train={np.mean(batch_losses):.4f}, "
            f"validation={validation_loss:.4f}",
            flush=True,
        )
        if fixed_epochs is None and stale >= 5:
            break
    model.load_state_dict(best_state)
    model.cpu().eval()
    return model, history, (mean, sd), best_epoch


def predict_adapter(
    model: TimesFMGraphAdapter,
    features: np.ndarray,
    scale: tuple[np.ndarray, np.ndarray],
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    mean, sd = scale
    values = torch.from_numpy(np.nan_to_num((features - mean) / sd).astype(np.float32))
    locations, scales, logits = [], [], []
    with torch.no_grad():
        for start in range(0, len(values), 8):
            output = model(values[start : start + 8])
            locations.append(output[0].numpy())
            scales.append(output[1].numpy())
            logits.append(output[2].numpy())
    return np.concatenate(locations), np.concatenate(scales), np.concatenate(logits)


def lognormal_outbreak_probability(
    location: np.ndarray,
    scale: np.ndarray,
    thresholds: np.ndarray,
) -> np.ndarray:
    """Analytic probability that the lognormal count exceeds its threshold."""
    z = (np.log1p(thresholds) - location) / np.maximum(scale, 1e-6)
    probability = 0.5 * torch.erfc(
        torch.from_numpy(z.astype(np.float32)) / np.sqrt(2.0)
    ).numpy()
    probability[~np.isfinite(thresholds)] = np.nan
    return probability.astype(np.float32)


def rows_for_window(
    label: str,
    units: tuple[str, ...],
    dates: pd.DatetimeIndex,
    target_dates: np.ndarray,
    targets: np.ndarray,
    population: np.ndarray,
    thresholds: np.ndarray,
    samples: np.ndarray,
    probabilities: np.ndarray,
    start_date: pd.Timestamp,
    end_date: pd.Timestamp,
    decision_thresholds: dict[int, float],
) -> tuple[pd.DataFrame, np.ndarray]:
    rows = []
    horizon_samples = []
    for horizon_index, horizon in enumerate(HORIZONS):
        mask = split_element_mask(target_dates, start_date, end_date)[:, horizon_index]
        horizon_samples.append(samples[mask, :, horizon_index])
        for issue_index in np.flatnonzero(mask):
            for node, unit in enumerate(units):
                values = samples[issue_index, node, horizon_index]
                rows.append(
                    {
                        "model": label,
                        "fcode": unit,
                        "issue_date": dates[issue_index],
                        "target_date": pd.Timestamp(target_dates[issue_index, horizon_index]),
                        "horizon": horizon,
                        "target_cases": targets[issue_index, node, horizon_index],
                        "population": population[issue_index, node, horizon_index],
                        "predicted_mean": float(values.mean()),
                        "predicted_median": float(np.median(values)),
                        "lower_50": float(np.quantile(values, 0.25)),
                        "upper_50": float(np.quantile(values, 0.75)),
                        "lower_80": float(np.quantile(values, 0.10)),
                        "upper_80": float(np.quantile(values, 0.90)),
                        "lower_95": float(np.quantile(values, 0.025)),
                        "upper_95": float(np.quantile(values, 0.975)),
                        "outbreak_threshold": thresholds[issue_index, node, horizon_index],
                        "outbreak_probability": probabilities[issue_index, node, horizon_index],
                        "outbreak_decision_threshold": decision_thresholds[horizon],
                    }
                )
    frame = pd.DataFrame(rows)
    cube = np.stack(horizon_samples, axis=2)
    return frame, cube


def select_alert_thresholds(
    targets: np.ndarray,
    thresholds: np.ndarray,
    probabilities: np.ndarray,
    target_dates: np.ndarray,
) -> tuple[dict[int, float], dict[str, list[dict[str, float]]]]:
    validation = split_element_mask(target_dates, VALIDATION_START, VALIDATION_END)
    selected, grids = {}, {}
    for horizon_index, horizon in enumerate(HORIZONS):
        mask = validation[:, horizon_index]
        observed = targets[mask, :, horizon_index].reshape(-1)
        threshold = thresholds[mask, :, horizon_index].reshape(-1)
        probability = probabilities[mask, :, horizon_index].reshape(-1)
        scores = []
        for cutoff in np.round(np.arange(0.05, 0.81, 0.01), 2):
            metrics = outbreak_classification_metrics(
                observed, threshold, probability, decision_threshold=float(cutoff)
            )
            scores.append(
                {
                    "threshold": float(cutoff),
                    "f1": float(metrics["outbreak_f1"]),
                    "precision": float(metrics["outbreak_precision"]),
                    "specificity": float(metrics["outbreak_specificity"]),
                    "sensitivity": float(metrics["outbreak_sensitivity"]),
                }
            )
        eligible = [row for row in scores if row["precision"] >= 0.50]
        winner = max(
            eligible or scores,
            key=lambda row: (
                row["f1"], row["sensitivity"], row["specificity"],
                -abs(row["threshold"] - 0.5),
            ),
        )
        selected[horizon] = winner["threshold"]
        grids[str(horizon)] = scores
    return selected, grids


def save_result(
    directory: Path,
    predictions: pd.DataFrame,
    samples: np.ndarray,
    decision_thresholds: dict[int, float],
    manifest: dict[str, object],
) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    predictions.to_csv(directory / "predictions.csv", index=False)
    np.savez_compressed(directory / "predictive_samples.npz", samples=samples)
    score_predictions(
        predictions, samples, decision_threshold=decision_thresholds
    ).to_csv(directory / "metrics.csv", index=False)
    (directory / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )


def display_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame[
        [
            "model", "mae", "rmse", "crps", "outbreak_brier",
            "outbreak_precision", "outbreak_specificity", "outbreak_sensitivity",
            "outbreak_f1",
        ]
    ].copy()
    output.columns = [
        "Model", "MAE", "RMSE", "CRPS", "Brier", "Precision", "Specificity",
        "Sensitivity", "F1",
    ]
    return output


def append_report_copy(output: Path) -> Path:
    shutil.copytree(CURRENT_REPORT, output, dirs_exist_ok=True)
    original = CURRENT_REPORT / "Final_Dengue_Forecasting_Report_2021_2025.docx"
    document = Document(original)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("11. TimesFM 3 and TimesFM–GNN comparison", level=1)
    document.add_paragraph(
        "This additional experiment evaluates the official Google Research TimesFM 3.0 PyTorch "
        "checkpoint and a spatial extension. It is appended only to this copied report; Sections 1–10 "
        "and the original report file remain unchanged. TimesFM 3 is a 330-million-parameter frozen "
        "foundation model. The no-GNN baseline forecasts each commune independently using dengue cases "
        "as the target and the seven selected climate variables as past-only covariates."
    )
    document.add_paragraph(
        "TimesFM–GNN uses the same frozen TimesFM forecasts and nine predictive quantiles as its temporal "
        "features. A 32-unit, four-head graph-attention adapter then exchanges information across the same "
        "5,766 directed candidate edges used by TCN-GNN. The adapter includes commune and horizon "
        "embeddings and produces a lognormal count distribution plus a focal-BCE outbreak head. Adapter "
        "epochs and alert cutoffs are selected only on 2019–2020; the adapter is then refitted on data "
        "through 2020 and tested on 2021–2025. TimesFM weights remain frozen throughout."
    )
    warning = document.add_paragraph()
    warning.add_run("License limitation. ").bold = True
    warning.add_run(
        "The TimesFM 3 pretrained weights use the TimesFM Non-Commercial License v1.0 and are restricted "
        "to non-commercial, non-production evaluation. Source code is Apache-2.0. Model: "
        "https://huggingface.co/google/timesfm-3.0-pytorch; implementation: "
        "https://github.com/google-research/timesfm."
    )

    existing = {
        "TCN-GNN Base": ROOT / "results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE/metrics.csv",
        "DeepST-INLA (proposed hybrid)": ROOT / "results/final_multiobjective_model_2021_2025/metrics.csv",
        **{
            label: RESULTS / key / "metrics.csv" for key, label in MODEL_LABELS.items()
        },
    }
    frames = []
    for label, path in existing.items():
        frame = pd.read_csv(path)
        frame["model"] = label
        frames.append(frame)
    full = pd.concat(frames, ignore_index=True)
    summary = full.groupby("model", as_index=False)[
        [
            "mae", "rmse", "crps", "outbreak_brier", "outbreak_precision",
            "outbreak_specificity", "outbreak_sensitivity", "outbreak_f1",
        ]
    ].mean()
    document.add_heading("11.1 Locked 2021–2025 results", level=2)
    add_table(document, display_metrics(summary))
    detailed = full.copy()
    detailed["model"] = detailed["model"] + " H" + detailed["horizon"].astype(str)
    add_table(document, display_metrics(detailed))

    period_sources = {
        "TCN-GNN Base": ROOT / "results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE",
        "DeepST-INLA (proposed hybrid)": ROOT / "results/final_multiobjective_model_2021_2025",
        **{label: RESULTS / key for key, label in MODEL_LABELS.items()},
    }
    restricted_by_horizon, restricted_summary = build_period_metrics(
        period_sources, output, "2023-01-01", "2025-12-01"
    )
    document.add_heading("11.2 Restricted 2023–2025 results", level=2)
    add_table(document, display_metrics(restricted_summary))
    restricted_detail = restricted_by_horizon.copy()
    restricted_detail["model"] = (
        restricted_detail["model"] + " H" + restricted_detail["horizon"].astype(str)
    )
    add_table(document, display_metrics(restricted_detail))

    document.add_heading("11.3 H3 visualization", level=2)
    figure_path = output / "Fig_H3_TimesFM_comparison_2021_2025.png"
    document.add_picture(str(figure_path), width=Inches(7.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph(
        "Figure 4. H3 comparison for the same four illustrative communes used in the original report."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8.5)

    timesfm = summary.set_index("model").loc["TimesFM 3 (no GNN)"]
    graph = summary.set_index("model").loc["TimesFM 3–GNN"]
    paragraph = document.add_paragraph()
    paragraph.add_run("Interpretation. ").bold = True
    paragraph.add_run(
        f"Adding the graph adapter changed mean H1–H3 MAE from {timesfm.mae:.3f} to "
        f"{graph.mae:.3f}, CRPS from {timesfm.crps:.3f} to {graph.crps:.3f}, Brier score from "
        f"{timesfm.outbreak_brier:.3f} to {graph.outbreak_brier:.3f}, and F1 from "
        f"{timesfm.outbreak_f1:.3f} to {graph.outbreak_f1:.3f}. Because TimesFM is zero-shot while "
        "the graph adapter is supervised, this contrast measures the combined value of dengue-specific "
        "adaptation and spatial graph learning; it is not a pure graph-only ablation."
    )
    path = output / "Final_Dengue_Forecasting_Report_2021_2025_TimesFM_Comparison.docx"
    document.save(path)
    return path


def build_figure(output: Path) -> None:
    command = [
        sys.executable,
        str(ROOT / "code/scripts/make_comparison_h3_outbreak_figure.py"),
        "--results-root", str(ROOT),
        "--output", str(output / "Fig_H3_TimesFM_comparison_2021_2025"),
        "--communes", "28159,27976,31147,31168",
        "--models",
        f"tcn={ROOT / 'results/feature_tcn_gnn_2021_2025_v1/ARCH-TCN-GNN-NOGATE/predictions.csv'}",
        f"timesfm={RESULTS / 'TIMESFM-3/predictions.csv'}",
        f"timesfm_gnn={RESULTS / 'TIMESFM-3-GNN/predictions.csv'}",
        f"hybrid={ROOT / 'results/final_multiobjective_model_2021_2025/predictions.csv'}",
        "--labels", "TCN-GNN Base", "TimesFM 3 (no GNN)", "TimesFM 3-GNN",
        "DeepST-INLA (proposed hybrid)",
        "--title", "Locked 2021-2025 H3 comparison: TCN-GNN and TimesFM alternatives",
    ]
    subprocess.run(command, check=True)


def main() -> None:
    global RESULTS
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", type=Path, default=RESULTS)
    parser.add_argument("--report-copy", type=Path, default=COPY_REPORT)
    parser.add_argument("--timesfm-batch-size", type=int, default=64)
    parser.add_argument("--device", default="cpu", choices=["cpu", "cuda"])
    parser.add_argument("--skip-report", action="store_true")
    parser.add_argument(
        "--pending-summary",
        type=Path,
        default=ROOT / "results/pending/timesfm3_summary.csv",
    )
    args = parser.parse_args()
    RESULTS = args.results.resolve()

    _, units, dates, cases, population, climate = load_panel()
    RESULTS.mkdir(parents=True, exist_ok=True)
    point, quantiles = build_backbone_cache(
        RESULTS / "timesfm3_backbone_cache.npz",
        units,
        dates,
        cases,
        climate,
        args.timesfm_batch_size,
        args.device,
    )
    thresholds = moh_thresholds(units, dates, cases)
    targets, target_population, target_thresholds, target_dates = target_arrays(
        dates, cases, population, thresholds
    )
    baseline_samples = quantiles_to_draws(quantiles)
    baseline_probability = np.nanmean(
        baseline_samples > target_thresholds[..., None], axis=-1
    ).astype(np.float32)
    baseline_cutoffs, baseline_threshold_grid = select_alert_thresholds(
        targets, target_thresholds, baseline_probability, target_dates
    )
    baseline_predictions, baseline_test_samples = rows_for_window(
        "TIMESFM-3", units, dates, target_dates, targets, target_population,
        target_thresholds, baseline_samples, baseline_probability, TEST_START,
        TEST_END, baseline_cutoffs,
    )
    save_result(
        RESULTS / "TIMESFM-3",
        baseline_predictions,
        baseline_test_samples,
        baseline_cutoffs,
        {
            "model": "TIMESFM-3",
            "checkpoint": CHECKPOINT,
            "timesfm_version": "3.0.1",
            "mode": "frozen zero-shot with past-only climate covariates; no graph",
            "climate_covariates": list(CLIMATE_COLUMNS),
            "quantiles_to_draws": "piecewise-linear interpolation of q0.1--q0.9 with linear tails",
            "decision_thresholds": baseline_cutoffs,
            "threshold_selection_grid": baseline_threshold_grid,
            "license": "timesfm-non-commercial-license-v1.0",
            "test": "2021-01 through 2025-12",
        },
    )

    features = backbone_features(point, quantiles, cases, population, dates)
    train = split_element_mask(target_dates, pd.Timestamp("2011-01-01"), TRAIN_END)
    validation = split_element_mask(target_dates, VALIDATION_START, VALIDATION_END)
    graph = build_candidate_graph(
        units, CENTROIDS_PATH, ADJACENCY_PATH, k_nearest=8,
        maximum_distance_km=100.0, include_administrative_edges=True,
    )
    selection_model, selection_history, selection_scale, best_epoch = fit_adapter(
        features,
        targets,
        target_thresholds,
        train,
        validation,
        graph,
        device_name=args.device,
    )
    selection_location, selection_scale_values, _ = predict_adapter(
        selection_model, features, selection_scale
    )
    selection_probability = lognormal_outbreak_probability(
        selection_location, selection_scale_values, target_thresholds
    )
    cutoffs, threshold_grid = select_alert_thresholds(
        targets, target_thresholds, selection_probability, target_dates
    )
    development = train | validation
    refitted, refit_history, refit_scale, _ = fit_adapter(
        features, targets, target_thresholds, development,
        np.zeros_like(validation), graph, fixed_epochs=best_epoch,
        device_name=args.device,
    )
    torch.save(
        {
            "state_dict": refitted.state_dict(),
            "feature_mean": refit_scale[0],
            "feature_sd": refit_scale[1],
            "selected_epoch": best_epoch,
        },
        RESULTS / "TIMESFM-3-GNN-model.pt",
    )
    pd.DataFrame(selection_history).to_csv(
        RESULTS / "TIMESFM-3-GNN-selection_history.csv", index=False
    )
    pd.DataFrame(refit_history).to_csv(
        RESULTS / "TIMESFM-3-GNN-refit_history.csv", index=False
    )
    location, scale, _ = predict_adapter(refitted, features, refit_scale)
    graph_samples = predictive_sample(
        location, scale, likelihood="lognormal", draws=500,
        rng=np.random.default_rng(1234),
    )
    graph_probability = lognormal_outbreak_probability(
        location, scale, target_thresholds
    )
    graph_predictions, graph_test_samples = rows_for_window(
        "TIMESFM-3-GNN", units, dates, target_dates, targets, target_population,
        target_thresholds, graph_samples, graph_probability, TEST_START, TEST_END,
        cutoffs,
    )
    save_result(
        RESULTS / "TIMESFM-3-GNN",
        graph_predictions,
        graph_test_samples,
        cutoffs,
        {
            "model": "TIMESFM-3-GNN",
            "checkpoint": CHECKPOINT,
            "timesfm_version": "3.0.1",
            "timesfm_weights": "frozen",
            "adapter": "32-unit four-head graph attention with commune/horizon embeddings",
            "candidate_edges_directed": int(graph.mask.sum()),
            "training_targets": "2011-01 through 2018-12",
            "validation_targets": "2019-01 through 2020-12",
            "refit_targets": "2011-01 through 2020-12",
            "selected_epoch": best_epoch,
            "decision_thresholds": cutoffs,
            "alert_probability_source": "fitted lognormal count distribution",
            "threshold_selection_grid": threshold_grid,
            "license": "timesfm-non-commercial-license-v1.0",
            "test": "2021-01 through 2025-12",
        },
    )
    summary_frames = []
    for key, label in MODEL_LABELS.items():
        frame = pd.read_csv(RESULTS / key / "metrics.csv")
        frame["model"] = label
        summary_frames.append(frame)
    summary = pd.concat(summary_frames, ignore_index=True)
    summary.insert(0, "checkpoint", CHECKPOINT)
    args.pending_summary.parent.mkdir(parents=True, exist_ok=True)
    summary.to_csv(args.pending_summary, index=False)
    print(f"Wrote {args.pending_summary}", flush=True)
    if not args.skip_report:
        build_figure(args.report_copy)
        report = append_report_copy(args.report_copy)
        print(f"Wrote {report}", flush=True)


if __name__ == "__main__":
    main()
